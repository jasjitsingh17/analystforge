"""Excel DCF workbook generator.

Architecture
------------
    DCFInputs -> run_dcf() -> DCFResult -> generate_dcf_workbook() -> .xlsx

The Python DCF engine (backend/dcf.py) remains the single source of truth for
the valuation. This module does NOT re-implement it. It does two separate
things with the ``DCFResult`` it receives:

1. It lays out the *inputs* (assumptions and historical data) as blue,
   hard-coded cells, and builds the model on top of them with visible Excel
   formulas (forecast, discounting, terminal value, EV bridge, equity bridge,
   sensitivity table). Change an assumption in Excel and the model updates.
2. It writes the engine's *results* into the 07_Model_Checks sheet as
   hard-coded reference values (purple), next to the values the Excel
   formulas calculate, with a difference and a PASS / FAIL flag. This is how
   any discrepancy between Excel and Python becomes visible.

No valuation output on the Cover, DCF, Sensitivity or Summary sheets is a
hard-coded number; they are all formulas.

Colour code used throughout the workbook
----------------------------------------
    Blue    hard-coded input        Black   formula on the same sheet
    Green   link to another sheet   Purple  Python engine reference value

Formulas are stored without cached values (openpyxl cannot calculate). Excel,
Numbers and LibreOffice calculate them when the file is opened; the workbook
is flagged to force a full calculation on load.
"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from backend.dcf import CONTROLLED_SAMPLE_HISTORICALS, DCFResult

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

SHEET_COVER = "00_Cover"
SHEET_ASSUMPTIONS = "01_Assumptions"
SHEET_HISTORICAL = "02_Historical"
SHEET_FORECAST = "03_Forecast"
SHEET_DCF = "04_DCF"
SHEET_SENSITIVITY = "05_Sensitivity"
SHEET_SUMMARY = "06_Valuation_Summary"
SHEET_CHECKS = "07_Model_Checks"
SHEET_SOURCES = "08_Data_Sources"

SHEET_NAMES: tuple[str, ...] = (
    SHEET_COVER,
    SHEET_ASSUMPTIONS,
    SHEET_HISTORICAL,
    SHEET_FORECAST,
    SHEET_DCF,
    SHEET_SENSITIVITY,
    SHEET_SUMMARY,
    SHEET_CHECKS,
    SHEET_SOURCES,
)

# Default output folder: <project root>/outputs, derived from this file's
# location at run time (no machine-specific absolute path is hard-coded).
DEFAULT_OUTPUT_DIR = Path(__file__).resolve().parent.parent / "outputs"

# Absolute tolerance for Python-vs-Excel reconciliation, in currency units.
# Double-precision noise on these numbers is ~1e-10, so 1e-6 is deliberately
# tight; it is visible (and editable) on the 07_Model_Checks sheet.
CHECK_TOLERANCE = 1e-6

# Sensitivity axis step (percentage points as decimals): 0.005 = 0.5%.
SENSITIVITY_STEP = 0.005

# Number formats. Zero renders as "-" and negatives in parentheses.
FMT_NUM = '#,##0;(#,##0);"-"'
FMT_PRICE = "#,##0.00;(#,##0.00);0.00"
FMT_PCT1 = "0.0%;(0.0%);0.0%"
FMT_PCT2 = "0.00%;(0.00%);0.00%"
FMT_FACTOR = "0.0000"
FMT_CHECK = "#,##0.000000;(#,##0.000000);0.000000"
FMT_YEARS = '0" years"'

_FONT = "Arial"
_NAVY = "1F3864"
_GREY = "F2F2F2"
_KEY_FILL = "DDEBF7"
_KIND_COLOR = {
    "input": "0000FF",
    "link": "008000",
    "calc": "000000",
    "label": "000000",
    "python": "7030A0",
}
_THIN = Side(style="thin", color="A6A6A6")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_BOTTOM = Border(bottom=Side(style="thin", color="000000"))
_TOP = Border(top=Side(style="thin", color="000000"))

# Row layout of 04_DCF (referenced by other sheets).
_D = {
    "period": 5, "ufcf": 6, "df": 7, "pv": 8,
    "wacc": 11, "g": 12, "final_ufcf": 13, "tv": 14, "df_final": 15, "pv_tv": 16,
    "sum_pv": 19, "pv_tv_bridge": 20, "ev": 21,
    "ev_bf": 24, "net_debt": 25, "equity": 26,
    "equity_bf": 29, "shares": 30, "price": 31, "cur_price": 32, "upside": 33,
    "tv_pct": 35,
}

# Row layout of 03_Forecast.
_F = {
    "period": 5, "revenue": 6, "growth": 7, "ebitda": 8, "margin": 9, "da": 10,
    "ebit": 11, "tax_rate": 12, "tax": 13, "nopat": 14, "capex": 15, "nwc": 16,
    "ufcf": 17,
}


# ---------------------------------------------------------------------------
# Small styling / addressing helpers
# ---------------------------------------------------------------------------


def _abs(sheet: str, cell: str) -> str:
    """Absolute cross-sheet reference, e.g. ('01_Assumptions','B9') -> '01_Assumptions'!$B$9."""
    col, row = coordinate_from_string(cell)
    return f"'{sheet}'!${col}${row}"


def _rel(sheet: str, cell: str) -> str:
    """Relative cross-sheet reference, e.g. '03_Forecast'!D17."""
    return f"'{sheet}'!{cell}"


def _range(sheet: str, start: str, end: str) -> str:
    """Cross-sheet range, e.g. '03_Forecast'!D5:H5."""
    return f"'{sheet}'!{start}:{end}"


def _put(
    ws: Worksheet,
    cell: str,
    value: object,
    *,
    kind: str = "calc",
    fmt: str | None = None,
    bold: bool = False,
    italic: bool = False,
    size: int = 10,
    fill: str | None = None,
    border: Border | None = None,
    align: str | None = None,
    wrap: bool = False,
) -> None:
    c = ws[cell]
    c.value = value
    c.font = Font(name=_FONT, size=size, bold=bold, italic=italic, color=_KIND_COLOR[kind])
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = PatternFill("solid", start_color=fill, end_color=fill)
    if border:
        c.border = border
    if align or wrap:
        c.alignment = Alignment(
            horizontal=align, vertical="top" if wrap else None, wrap_text=wrap
        )


def _title(ws: Worksheet, title: str, subtitle: str) -> None:
    ws.sheet_view.showGridLines = False
    _put(ws, "A1", title, kind="label", bold=True, size=14)
    ws["A1"].font = Font(name=_FONT, size=14, bold=True, color=_NAVY)
    _put(ws, "A2", subtitle, kind="label", italic=True)
    ws["A2"].font = Font(name=_FONT, size=9, italic=True, color="595959")
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def _section(ws: Worksheet, row: int, text: str, last_col: int) -> None:
    for col in range(1, last_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = PatternFill("solid", start_color=_NAVY, end_color=_NAVY)
        cell.font = Font(name=_FONT, size=10, bold=True, color="FFFFFF")
    ws.cell(row=row, column=1).value = text


def _header(
    ws: Worksheet,
    row: int,
    labels: list[str],
    first_col: int = 1,
    left: tuple[int, ...] = (0,),
) -> None:
    """Write a table header row; columns whose index is in ``left`` are left-aligned."""
    for offset, label in enumerate(labels):
        cell = ws.cell(row=row, column=first_col + offset)
        cell.value = label
        cell.font = Font(name=_FONT, size=10, bold=True, color="000000")
        cell.fill = PatternFill("solid", start_color=_GREY, end_color=_GREY)
        cell.border = _BOTTOM
        cell.alignment = Alignment(horizontal="left" if offset in left else "right")


def _widths(ws: Worksheet, widths: dict[str, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


# ---------------------------------------------------------------------------
# 01_Assumptions
# ---------------------------------------------------------------------------


def _build_assumptions(ws: Worksheet, result: DCFResult) -> dict[str, str]:
    """Write the assumptions table; return absolute references keyed by field."""
    inp = result.inputs
    _title(
        ws,
        "Model Assumptions",
        "Blue = hard-coded input | Black = formula | Green = link to another sheet. "
        "All monetary amounts share one unspecified currency unit.",
    )
    _header(ws, 4, ["Assumption", "Value", "Unit", "Notes"], left=(0, 2, 3))
    spec = [
        ("company_name", "Company name", inp.company_name, None, "text", "Company being valued"),
        ("ticker", "Ticker", inp.ticker, None, "text", "Identifier used for labelling only"),
        ("share_price", "Current share price", inp.share_price, FMT_PRICE, "currency / share",
         "Supplied as an input; not a live market quote"),
        ("forecast_years", "Forecast years", inp.forecast_years, "0", "years",
         "Number of forecast columns is fixed when the workbook is generated"),
        ("revenue_growth", "Revenue growth", inp.revenue_growth, FMT_PCT2, "% per year",
         "Constant annual growth applied to every forecast year"),
        ("ebitda_margin", "EBITDA margin", inp.ebitda_margin, FMT_PCT2, "% of revenue",
         "Constant margin applied to every forecast year"),
        ("tax_rate", "Tax rate", inp.tax_rate, FMT_PCT2, "% of EBIT", "Applied to EBIT (unlevered tax)"),
        ("wacc", "WACC", inp.wacc, FMT_PCT2, "% per year", "Discount rate for cash flows and terminal value"),
        ("terminal_growth", "Terminal growth", inp.terminal_growth, FMT_PCT2, "% per year",
         "Perpetual growth rate after the final forecast year"),
        ("net_debt", "Net debt", inp.net_debt, FMT_NUM, "currency units",
         "Debt less cash; deducted from enterprise value"),
        ("shares_outstanding", "Shares outstanding", inp.shares_outstanding, FMT_NUM, "shares",
         "Must be on the same unit scale as the monetary amounts"),
    ]
    refs: dict[str, str] = {}
    for i, (key, label, value, fmt, unit, note) in enumerate(spec):
        row = 5 + i
        _put(ws, f"A{row}", label, kind="label")
        _put(ws, f"B{row}", value, kind="input", fmt=fmt, align="right", border=_BOX)
        _put(ws, f"C{row}", unit, kind="label")
        _put(ws, f"D{row}", note, kind="label", italic=True)
        refs[key] = _abs(ws.title, f"B{row}")
    _widths(ws, {"A": 26, "B": 22, "C": 20, "D": 70})
    ws.freeze_panes = "A5"
    return refs


# ---------------------------------------------------------------------------
# 02_Historical
# ---------------------------------------------------------------------------


def _build_historical(ws: Worksheet, result: DCFResult) -> dict[str, str | None]:
    """Write the historical data used by the engine; return references."""
    hist = result.historical
    n = len(hist.revenue)
    is_sample = hist == CONTROLLED_SAMPLE_HISTORICALS
    _title(
        ws,
        "Historical Financials",
        "Controlled sample data used for demonstration and testing - not live or "
        "company-reported data."
        if is_sample
        else "Custom historical data supplied by the caller - not verified by this model.",
    )
    labels = ["Line item"]
    for i in range(n):
        tag = " (latest)" if i == n - 1 else " (oldest)" if i == 0 else ""
        labels.append(f"Period {i + 1}{tag}")
    _header(ws, 4, labels)

    cols = [get_column_letter(2 + i) for i in range(n)]
    series = [
        (5, "Revenue", hist.revenue),
        (7, "EBITDA", hist.ebitda),
        (9, "D&A", hist.depreciation_amortization),
        (10, "CapEx", hist.capex),
        (11, "Change in NWC", hist.change_in_nwc),
    ]
    for row, label, values in series:
        _put(ws, f"A{row}", label, kind="label")
        for col, value in zip(cols, values):
            _put(ws, f"{col}{row}", value, kind="input", fmt=FMT_NUM)
    _put(ws, "A6", "Revenue growth", kind="label", italic=True)
    _put(ws, "A8", "EBITDA margin", kind="label", italic=True)
    for i, col in enumerate(cols):
        if i > 0:
            _put(ws, f"{col}6", f"={col}5/{cols[i - 1]}5-1", fmt=FMT_PCT1, italic=True)
        _put(ws, f"{col}8", f"={col}7/{col}5", fmt=FMT_PCT1, italic=True)

    last = cols[-1]
    _section(ws, 13, "Forecast base values (latest historical period)", n + 1)
    base = [
        (14, "Latest revenue", f"={last}5", "latest_revenue"),
        (15, "Latest D&A", f"={last}9", "latest_da"),
        (16, "Latest CapEx", f"={last}10", "latest_capex"),
        (17, "Latest change in NWC", f"={last}11", "latest_nwc"),
    ]
    refs: dict[str, str | None] = {}
    for row, label, formula, key in base:
        _put(ws, f"A{row}", label, kind="label")
        _put(ws, f"B{row}", formula, fmt=FMT_NUM, border=_BOX)
        refs[key] = _abs(ws.title, f"B{row}")
    refs["latest_ebitda"] = _abs(ws.title, f"{last}7")
    refs["latest_margin"] = _abs(ws.title, f"{last}8")
    refs["latest_growth"] = _abs(ws.title, f"{last}6") if n >= 2 else None

    notes = [
        "The DCF engine forecasts D&A, CapEx and Change in NWC by holding their latest "
        "historical value flat (a version 1 simplification).",
        "Historical EBITDA is shown for reference only: the engine forecasts EBITDA from "
        "the EBITDA margin assumption, not from history.",
        "Monetary amounts share one unspecified currency unit.",
    ]
    for i, note in enumerate(notes):
        _put(ws, f"A{19 + i}", note, kind="label", italic=True)
    _widths(ws, {"A": 26, **{c: 16 for c in cols}})
    ws.freeze_panes = "B5"
    return refs


# ---------------------------------------------------------------------------
# 03_Forecast
# ---------------------------------------------------------------------------


def _build_forecast(
    ws: Worksheet, result: DCFResult, a: dict[str, str], h: dict[str, str | None]
) -> dict[str, object]:
    """Write the forecast table with live formulas; return column info."""
    n = len(result.projections)
    cols = [get_column_letter(4 + i) for i in range(n)]  # D, E, ...
    _title(
        ws,
        "Forecast",
        "Every forecast cell is a formula linked to 01_Assumptions and 02_Historical. "
        "Column C is the latest historical period (base year).",
    )
    _header(ws, 4, ["Line item", "Basis", "Base (latest hist.)"] + [f"Year {i + 1}" for i in range(n)],
            left=(0, 1))

    rows = [
        ("period", "Forecast period (t)", "Year index used for discounting", "0"),
        ("revenue", "Revenue", "Prior year x (1 + revenue growth)", FMT_NUM),
        ("growth", "Revenue growth", "Assumption (constant)", FMT_PCT1),
        ("ebitda", "EBITDA", "Revenue x EBITDA margin", FMT_NUM),
        ("margin", "EBITDA margin", "Assumption (constant)", FMT_PCT1),
        ("da", "D&A", "Held flat at latest historical value", FMT_NUM),
        ("ebit", "EBIT", "EBITDA - D&A", FMT_NUM),
        ("tax_rate", "Tax rate", "Assumption", FMT_PCT1),
        ("tax", "Tax", "EBIT x tax rate", FMT_NUM),
        ("nopat", "NOPAT", "EBIT - tax", FMT_NUM),
        ("capex", "CapEx", "Held flat at latest historical value", FMT_NUM),
        ("nwc", "Change in NWC", "Held flat at latest historical value", FMT_NUM),
        ("ufcf", "UFCF", "NOPAT + D&A - CapEx - change in NWC", FMT_NUM),
    ]
    for key, label, basis, fmt in rows:
        r = _F[key]
        bold = key in {"revenue", "ebitda", "ebit", "nopat", "ufcf"}
        _put(ws, f"A{r}", label, kind="label", bold=bold)
        _put(ws, f"B{r}", basis, kind="label", italic=True)
        if key == "ufcf":
            for col in ["A", "B", "C"] + cols:
                ws[f"{col}{r}"].border = _TOP
    # Base year column (C): links to the historical sheet, no calculations.
    _put(ws, "C5", 0, kind="input", fmt="0")
    base_links = [
        ("revenue", h["latest_revenue"], FMT_NUM),
        ("growth", h["latest_growth"], FMT_PCT1),
        ("ebitda", h["latest_ebitda"], FMT_NUM),
        ("margin", h["latest_margin"], FMT_PCT1),
        ("da", h["latest_da"], FMT_NUM),
        ("capex", h["latest_capex"], FMT_NUM),
        ("nwc", h["latest_nwc"], FMT_NUM),
    ]
    for key, ref, fmt in base_links:
        if ref:
            _put(ws, f"C{_F[key]}", f"={ref}", kind="link", fmt=fmt)

    for i, col in enumerate(cols):
        prev = "C" if i == 0 else cols[i - 1]
        f = _F
        cells = [
            ("period", f"={prev}{f['period']}+1", "calc", "0"),
            ("revenue", f"={prev}{f['revenue']}*(1+{col}{f['growth']})", "calc", FMT_NUM),
            ("growth", f"={a['revenue_growth']}", "link", FMT_PCT1),
            ("ebitda", f"={col}{f['revenue']}*{col}{f['margin']}", "calc", FMT_NUM),
            ("margin", f"={a['ebitda_margin']}", "link", FMT_PCT1),
            ("da", f"={prev}{f['da']}", "calc", FMT_NUM),
            ("ebit", f"={col}{f['ebitda']}-{col}{f['da']}", "calc", FMT_NUM),
            ("tax_rate", f"={a['tax_rate']}", "link", FMT_PCT1),
            ("tax", f"={col}{f['ebit']}*{col}{f['tax_rate']}", "calc", FMT_NUM),
            ("nopat", f"={col}{f['ebit']}-{col}{f['tax']}", "calc", FMT_NUM),
            ("capex", f"={prev}{f['capex']}", "calc", FMT_NUM),
            ("nwc", f"={prev}{f['nwc']}", "calc", FMT_NUM),
            (
                "ufcf",
                f"={col}{f['nopat']}+{col}{f['da']}-{col}{f['capex']}-{col}{f['nwc']}",
                "calc",
                FMT_NUM,
            ),
        ]
        for key, formula, kind, fmt in cells:
            _put(
                ws, f"{col}{f[key]}", formula, kind=kind, fmt=fmt,
                bold=key in {"revenue", "ebitda", "ebit", "nopat", "ufcf"},
                border=_TOP if key == "ufcf" else None,
            )
    _put(ws, "A19", "CapEx and Change in NWC are shown as positive amounts and subtracted "
                    "in the UFCF formula (same convention as the Python engine).",
         kind="label", italic=True)
    _widths(ws, {"A": 26, "B": 40, "C": 20, **{c: 14 for c in cols}})
    ws.freeze_panes = "C6"
    return {"cols": cols, "first": cols[0], "last": cols[-1]}


# ---------------------------------------------------------------------------
# 04_DCF
# ---------------------------------------------------------------------------


def _build_dcf(
    ws: Worksheet, result: DCFResult, a: dict[str, str], fc: dict[str, object]
) -> dict[str, str]:
    """Write discounting, terminal value and the EV / equity / price bridges."""
    n = len(result.projections)
    cols = [get_column_letter(3 + i) for i in range(n)]  # C, D, ...
    fcols: list[str] = fc["cols"]  # type: ignore[assignment]
    last = cols[-1]
    d = _D
    _title(
        ws,
        "DCF Valuation",
        "End-of-year discounting. All cells are formulas linked to 03_Forecast and 01_Assumptions.",
    )
    _header(ws, 4, ["Discounted cash flows", "Basis"] + [f"Year {i + 1}" for i in range(n)],
            left=(0, 1))
    labels = [
        ("period", "Forecast period (t)", "Year index"),
        ("ufcf", "UFCF", "From 03_Forecast"),
        ("df", "Discount factor", "1 / (1 + WACC) ^ t"),
        ("pv", "PV of UFCF", "UFCF x discount factor"),
    ]
    for key, label, basis in labels:
        _put(ws, f"A{d[key]}", label, kind="label", bold=key == "pv")
        _put(ws, f"B{d[key]}", basis, kind="label", italic=True)
    for col, fcol in zip(cols, fcols):
        period_link = _rel(SHEET_FORECAST, f"{fcol}{_F['period']}")
        ufcf_link = _rel(SHEET_FORECAST, f"{fcol}{_F['ufcf']}")
        _put(ws, f"{col}{d['period']}", f"={period_link}", kind="link", fmt="0")
        _put(ws, f"{col}{d['ufcf']}", f"={ufcf_link}", kind="link", fmt=FMT_NUM)
        _put(ws, f"{col}{d['df']}", f"=1/(1+$C${d['wacc']})^{col}{d['period']}", fmt=FMT_FACTOR)
        _put(ws, f"{col}{d['pv']}", f"={col}{d['ufcf']}*{col}{d['df']}", fmt=FMT_NUM, bold=True,
             border=_TOP)

    def line(key: str, label: str, basis: str, formula: str, fmt: str, *, kind: str = "calc",
             bold: bool = False, key_output: bool = False) -> None:
        r = d[key]
        _put(ws, f"A{r}", label, kind="label", bold=bold or key_output)
        _put(ws, f"B{r}", basis, kind="label", italic=True)
        _put(ws, f"C{r}", formula, kind=kind, fmt=fmt, bold=bold or key_output,
             fill=_KEY_FILL if key_output else None, border=_BOX if key_output else None)

    _section(ws, 10, "Terminal value (Gordon growth)", 2 + n)
    line("wacc", "WACC", "Assumption", f"={a['wacc']}", FMT_PCT2, kind="link")
    line("g", "Terminal growth", "Assumption", f"={a['terminal_growth']}", FMT_PCT2, kind="link")
    line("final_ufcf", "Final year UFCF", "UFCF in last forecast year", f"={last}{d['ufcf']}", FMT_NUM)
    line("tv", "Terminal value", "Final UFCF x (1 + g) / (WACC - g)",
         f"=C{d['final_ufcf']}*(1+C{d['g']})/(C{d['wacc']}-C{d['g']})", FMT_NUM, bold=True)
    line("df_final", "Discount factor (final year)", "Discount factor of last forecast year",
         f"={last}{d['df']}", FMT_FACTOR)
    line("pv_tv", "PV of terminal value (calculation)", "Terminal value x final-year discount factor",
         f"=C{d['tv']}*C{d['df_final']}", FMT_NUM, bold=True)

    _section(ws, 18, "Enterprise value bridge", 2 + n)
    line("sum_pv", "PV of forecast UFCF", "Sum of PV of UFCF (row 8)",
         f"=SUM(C{d['pv']}:{last}{d['pv']})", FMT_NUM)
    line("pv_tv_bridge", "PV of terminal value", "Plus: from terminal value section",
         f"=C{d['pv_tv']}", FMT_NUM)
    line("ev", "Enterprise value", "PV of forecast UFCF + PV of terminal value",
         f"=C{d['sum_pv']}+C{d['pv_tv_bridge']}", FMT_NUM, key_output=True)

    _section(ws, 23, "Equity value bridge", 2 + n)
    line("ev_bf", "Enterprise value (brought forward)", "From enterprise value bridge",
         f"=C{d['ev']}", FMT_NUM)
    line("net_debt", "Net debt", "Less: assumption", f"={a['net_debt']}", FMT_NUM, kind="link")
    line("equity", "Equity value", "Enterprise value - net debt",
         f"=C{d['ev_bf']}-C{d['net_debt']}", FMT_NUM, key_output=True)

    _section(ws, 28, "Per-share valuation", 2 + n)
    line("equity_bf", "Equity value (brought forward)", "From equity value bridge",
         f"=C{d['equity']}", FMT_NUM)
    line("shares", "Shares outstanding", "Assumption", f"={a['shares_outstanding']}", FMT_NUM,
         kind="link")
    line("price", "Implied share price", "Equity value / shares outstanding",
         f"=C{d['equity_bf']}/C{d['shares']}", FMT_PRICE, key_output=True)
    line("cur_price", "Current share price", "Assumption", f"={a['share_price']}", FMT_PRICE,
         kind="link")
    line("upside", "Upside / (downside)", "(Implied - current) / current",
         f"=(C{d['price']}-C{d['cur_price']})/C{d['cur_price']}", FMT_PCT1, key_output=True)

    line("tv_pct", "Memo: PV of terminal value as % of EV", "PV of terminal value / enterprise value",
         f"=C{d['pv_tv_bridge']}/C{d['ev']}", FMT_PCT1)
    ws[f"A{d['tv_pct']}"].font = Font(name=_FONT, size=10, italic=True)

    _widths(ws, {"A": 38, "B": 44, **{c: 14 for c in cols}})
    ws.freeze_panes = "C5"
    return {
        "first": cols[0],
        "last": last,
        "ev": _abs(SHEET_DCF, f"C{d['ev']}"),
        "equity": _abs(SHEET_DCF, f"C{d['equity']}"),
        "price": _abs(SHEET_DCF, f"C{d['price']}"),
        "upside": _abs(SHEET_DCF, f"C{d['upside']}"),
        "sum_pv": _abs(SHEET_DCF, f"C{d['sum_pv']}"),
        "tv": _abs(SHEET_DCF, f"C{d['tv']}"),
        "pv_tv": _abs(SHEET_DCF, f"C{d['pv_tv']}"),
        "pv_tv_bridge": _abs(SHEET_DCF, f"C{d['pv_tv_bridge']}"),
        "net_debt": _abs(SHEET_DCF, f"C{d['net_debt']}"),
        "shares": _abs(SHEET_DCF, f"C{d['shares']}"),
    }


# ---------------------------------------------------------------------------
# 05_Sensitivity
# ---------------------------------------------------------------------------


def _build_sensitivity(
    ws: Worksheet, result: DCFResult, a: dict[str, str], dcf: dict[str, str]
) -> dict[str, object]:
    """Implied share price for a 5 x 5 grid of WACC and terminal growth.

    Because UFCF does not depend on WACC or terminal growth, each grid cell
    can be calculated directly with visible formulas: a helper block below the
    grid discounts every UFCF at each WACC, and each grid cell combines that
    with the Gordon-growth terminal value for its (WACC, g) pair. No Excel
    data table is used, so it works in every spreadsheet program.
    """
    n = len(result.projections)
    d = _D
    ycols = [get_column_letter(3 + i) for i in range(n)]  # helper year columns C..
    sum_col = get_column_letter(3 + n)
    dcf_last = str(dcf["last"])
    _title(
        ws,
        "Sensitivity Analysis: Implied Share Price",
        "Rows = WACC, columns = terminal growth. Centre cell (highlighted) is the base case. "
        "Formula-driven; n/a where WACC <= terminal growth.",
    )
    _put(ws, "A3", "WACC step", kind="label")
    _put(ws, "B3", SENSITIVITY_STEP, kind="input", fmt=FMT_PCT2, border=_BOX)
    _put(ws, "A4", "Terminal growth step", kind="label")
    _put(ws, "B4", SENSITIVITY_STEP, kind="input", fmt=FMT_PCT2, border=_BOX)

    _section(ws, 6, "Implied share price (currency per share)", 7)
    _put(ws, "A7", "WACC (rows) / Terminal growth (columns)", kind="label", bold=True)
    grid_cols = ["C", "D", "E", "F", "G"]
    grid_rows = [8, 9, 10, 11, 12]
    # Axes: the centre links to the base case; the others step away from it.
    # ROUND(..., 10) removes binary floating-point noise (e.g. 0.015-0.005 =
    # 0.010000000000000002) so equal rates compare as equal and the
    # WACC <= terminal growth guard behaves exactly.
    _put(ws, "E7", f"={a['terminal_growth']}", kind="link", fmt=FMT_PCT2, bold=True, fill=_GREY,
         border=_BOX)
    for col, k in zip(grid_cols, [-2, -1, 0, 1, 2]):
        if k != 0:
            _put(ws, f"{col}7", f"=ROUND($E$7{'+' if k > 0 else '-'}{abs(k)}*$B$4,10)",
                 fmt=FMT_PCT2, bold=True, fill=_GREY, border=_BOX)
    _put(ws, "B10", f"={a['wacc']}", kind="link", fmt=FMT_PCT2, bold=True, fill=_GREY, border=_BOX)
    for row, k in zip(grid_rows, [-2, -1, 0, 1, 2]):
        if k != 0:
            _put(ws, f"B{row}", f"=ROUND($B$10{'+' if k > 0 else '-'}{abs(k)}*$B$3,10)",
                 fmt=FMT_PCT2, bold=True, fill=_GREY, border=_BOX)

    ufcf_final = _abs(SHEET_DCF, f"{dcf_last}{d['ufcf']}")
    period_final = _abs(SHEET_DCF, f"{dcf_last}{d['period']}")
    for row in grid_rows:
        helper_row = row + 9  # 17..21
        for col in grid_cols:
            formula = (
                f'=IF(AND($B{row}>0,$B{row}>{col}$7),'
                f'(${sum_col}{helper_row}+{ufcf_final}*(1+{col}$7)/($B{row}-{col}$7)'
                f'/(1+$B{row})^{period_final}-{dcf["net_debt"]})/{dcf["shares"]},"n/a")'
            )
            _put(ws, f"{col}{row}", formula, fmt=FMT_PRICE, border=_BOX, align="right")
    ws["E10"].fill = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
    ws["E10"].font = Font(name=_FONT, size=10, bold=True)

    _section(ws, 14, "Calculation helper: PV of each year's UFCF at each WACC", 3 + n)
    _put(ws, "B15", "Forecast period (t)", kind="label", bold=True)
    _put(ws, "B16", "UFCF", kind="label", bold=True)
    for col in ycols:
        period_link = _rel(SHEET_DCF, f"{col}{d['period']}")
        ufcf_link = _rel(SHEET_DCF, f"{col}{d['ufcf']}")
        _put(ws, f"{col}15", f"={period_link}", kind="link", fmt="0")
        _put(ws, f"{col}16", f"={ufcf_link}", kind="link", fmt=FMT_NUM)
    _put(ws, f"{sum_col}15", "Sum of PV", kind="label", bold=True, align="right")
    for row in grid_rows:
        hr = row + 9
        _put(ws, f"A{hr}", "WACC", kind="label")
        _put(ws, f"B{hr}", f"=$B{row}", fmt=FMT_PCT2)
        for col in ycols:
            _put(ws, f"{col}{hr}", f'=IF($B{hr}<=0,"n/a",{col}$16/(1+$B{hr})^{col}$15)',
                 fmt=FMT_NUM, align="right")
        _put(ws, f"{sum_col}{hr}", f'=IF($B{hr}<=0,"n/a",SUM({ycols[0]}{hr}:{ycols[-1]}{hr}))',
             fmt=FMT_NUM, bold=True, align="right")
    _put(ws, "A23",
         "Each grid cell = (sum of PV of UFCF at that WACC + PV of terminal value at that WACC "
         "and growth - net debt) / shares outstanding.", kind="label", italic=True)
    _widths(ws, {"A": 40, "B": 22, **{c: 13 for c in ycols}, sum_col: 14})
    for col in grid_cols:
        ws.column_dimensions[col].width = max(ws.column_dimensions[col].width or 0, 13)
    return {"grid_cols": grid_cols, "grid_rows": grid_rows}


# ---------------------------------------------------------------------------
# 06_Valuation_Summary
# ---------------------------------------------------------------------------


def _build_summary(
    ws: Worksheet, a: dict[str, str], dcf: dict[str, str]
) -> None:
    _title(ws, "Valuation Summary", "Illustrative DCF output. All values link to the model sheets.")
    _header(ws, 4, ["Metric", "Value"])
    rows = [
        ("Company", f"={a['company_name']}", None, False),
        ("Ticker", f"={a['ticker']}", None, False),
        ("Current share price", f"={a['share_price']}", FMT_PRICE, False),
        ("Implied DCF share price", f"={dcf['price']}", FMT_PRICE, True),
        ("Upside / (downside)", f"={dcf['upside']}", FMT_PCT1, True),
        ("Enterprise value", f"={dcf['ev']}", FMT_NUM, True),
        ("Equity value", f"={dcf['equity']}", FMT_NUM, True),
        ("WACC", f"={a['wacc']}", FMT_PCT2, False),
        ("Terminal growth", f"={a['terminal_growth']}", FMT_PCT2, False),
    ]
    for i, (label, formula, fmt, key) in enumerate(rows):
        r = 5 + i
        _put(ws, f"A{r}", label, kind="label", bold=key, size=11 if key else 10)
        _put(ws, f"B{r}", formula, kind="link", fmt=fmt, bold=key, size=11 if key else 10,
             fill=_KEY_FILL if key else None, border=_BOX, align="right")
    _put(ws, "A15", "Sensitivity of the implied share price to WACC and terminal growth is on "
                    "05_Sensitivity; reconciliation to the Python engine is on 07_Model_Checks.",
         kind="label", italic=True)
    _widths(ws, {"A": 32, "B": 26})


# ---------------------------------------------------------------------------
# 07_Model_Checks
# ---------------------------------------------------------------------------


def _build_checks(
    ws: Worksheet,
    result: DCFResult,
    a: dict[str, str],
    dcf: dict[str, str],
    fc: dict[str, object],
) -> None:
    """Reconcile Python engine values with the values Excel calculates."""
    d = _D
    f = _F
    n = len(result.projections)
    fcols: list[str] = fc["cols"]  # type: ignore[assignment]
    period_range = _range(
        SHEET_FORECAST, f"{fcols[0]}{f['period']}", f"{fcols[-1]}{f['period']}"
    )
    ufcf_range = _range(SHEET_FORECAST, f"{fcols[0]}{f['ufcf']}", f"{fcols[-1]}{f['ufcf']}")
    dcf_pv_range = _range(SHEET_DCF, f"{dcf['first']}{d['pv']}", f"{dcf['last']}{d['pv']}")

    _title(
        ws,
        "Model Checks",
        "Python engine values (purple, hard-coded reference) vs values calculated by the Excel "
        "formulas in this workbook.",
    )
    first_row = 8
    # (name, reference value/formula, excel formula, mode, description)
    checks: list[tuple[str, object, str, str, str]] = [
        ("Python vs Excel: Enterprise value", float(result.enterprise_value), f"={dcf['ev']}",
         "money", "Enterprise value from the Python engine vs the 04_DCF formulas"),
        ("Python vs Excel: Equity value", float(result.equity_value), f"={dcf['equity']}",
         "money", "Equity value from the Python engine vs the 04_DCF formulas"),
        ("Python vs Excel: Implied share price", float(result.implied_share_price),
         f"={dcf['price']}", "money", "Implied share price from the Python engine vs the 04_DCF formulas"),
        ("Forecast year count", n, f"=COUNT({period_range})", "exact",
         "Number of forecast years in the Python result vs forecast columns in 03_Forecast"),
        ("Enterprise value bridge",
         f"=SUM({dcf_pv_range})+{dcf['pv_tv']}", f"={dcf['ev']}", "money",
         "Sum of yearly PVs + PV of terminal value (recomputed here) vs enterprise value shown"),
        ("Equity value bridge", f"={dcf['ev']}-{a['net_debt']}", f"={dcf['equity']}", "money",
         "Enterprise value - net debt (recomputed here) vs equity value shown"),
        ("Implied share price reconciliation", f"={dcf['equity']}/{a['shares_outstanding']}",
         f"={dcf['price']}", "money",
         "Equity value / shares outstanding (recomputed here) vs implied share price shown"),
        ("WACC greater than terminal growth", f"={a['wacc']}", f"={a['terminal_growth']}",
         "spread", "WACC (column C) must exceed terminal growth (column D); difference is the spread"),
        ("Python vs Excel: Sum of PV of UFCF", float(result.sum_pv_ufcf), f"={dcf['sum_pv']}",
         "money", "Sum of discounted forecast cash flows"),
        ("Python vs Excel: Terminal value", float(result.terminal_value), f"={dcf['tv']}",
         "money", "Undiscounted terminal value"),
        ("Python vs Excel: PV of terminal value", float(result.pv_terminal_value),
         f"={dcf['pv_tv']}", "money", "Terminal value discounted to today"),
        ("Python vs Excel: Total UFCF", float(sum(p.ufcf for p in result.projections)),
         f"=SUM({ufcf_range})", "money", "Sum of UFCF across all forecast years"),
        ("Forecast years assumption vs forecast columns", f"={a['forecast_years']}",
         f"=COUNT({period_range})", "exact",
         "Forecast years assumption must equal the number of forecast columns generated"),
    ]
    last_row = first_row + len(checks) - 1

    _put(ws, "A4", "Overall model status", kind="label", bold=True)
    _put(ws, "B4",
         f'=IF(COUNTIF(G{first_row}:G{last_row},"FAIL")=0,"ALL CHECKS PASS","CHECK FAILURE - REVIEW MODEL")',
         bold=True, border=_BOX)
    _put(ws, "A5", "Absolute tolerance (currency units)", kind="label", bold=True)
    _put(ws, "B5", CHECK_TOLERANCE, kind="input", fmt="0.000000", border=_BOX)
    _put(ws, "C5", "Tight on purpose: floating-point noise on these values is far smaller.",
         kind="label", italic=True)
    _header(ws, 7, ["#", "Check", "Python / reference", "Excel value", "Difference",
                    "Tolerance", "Result", "Description"], left=(0, 1, 7))
    for i, (name, ref, excel, mode, desc) in enumerate(checks):
        r = first_row + i
        is_formula = isinstance(ref, str)
        _put(ws, f"A{r}", i + 1, kind="label", align="left")
        _put(ws, f"B{r}", name, kind="label")
        if mode == "exact":
            ref_fmt, tol = "0", 0
        else:
            ref_fmt, tol = FMT_CHECK, "=$B$5"
        _put(ws, f"C{r}", ref, kind="calc" if is_formula else "python", fmt=ref_fmt)
        _put(ws, f"D{r}", excel, kind="link", fmt=ref_fmt)
        if mode == "spread":
            _put(ws, f"E{r}", f"=C{r}-D{r}", fmt=FMT_PCT2)
            _put(ws, f"F{r}", "spread > 0", kind="label", align="right")
            _put(ws, f"G{r}", f'=IF(E{r}>0,"PASS","FAIL")', bold=True, align="center")
            ws[f"C{r}"].number_format = FMT_PCT2
            ws[f"D{r}"].number_format = FMT_PCT2
        else:
            _put(ws, f"E{r}", f"=D{r}-C{r}", fmt=FMT_CHECK if mode == "money" else "0")
            _put(ws, f"F{r}", tol, fmt="0.000000" if mode == "money" else "0")
            _put(ws, f"G{r}", f'=IF(ABS(E{r})<=F{r},"PASS","FAIL")', bold=True, align="center")
        _put(ws, f"H{r}", desc, kind="label", italic=True)

    green = PatternFill("solid", start_color="C6EFCE", end_color="C6EFCE")
    red = PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")
    rng = f"G{first_row}:G{last_row}"
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=[f'G{first_row}="PASS"'], fill=green, font=Font(color="006100", bold=True)))
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=[f'G{first_row}="FAIL"'], fill=red, font=Font(color="9C0006", bold=True)))
    ws.conditional_formatting.add(
        "B4", FormulaRule(formula=['$B$4="ALL CHECKS PASS"'], fill=green, font=Font(color="006100", bold=True)))
    ws.conditional_formatting.add(
        "B4", FormulaRule(formula=['$B$4<>"ALL CHECKS PASS"'], fill=red, font=Font(color="9C0006", bold=True)))
    _put(ws, f"A{last_row + 2}",
         "Purple = hard-coded value taken from the tested Python DCF engine (backend/dcf.py). "
         "If a check fails, the Excel formulas and the Python engine disagree: investigate, do not "
         "widen the tolerance.", kind="label", italic=True)
    _widths(ws, {"A": 34, "B": 46, "C": 22, "D": 22, "E": 18, "F": 14, "G": 10, "H": 80})
    ws.freeze_panes = "A8"


# ---------------------------------------------------------------------------
# 08_Data_Sources
# ---------------------------------------------------------------------------


def _build_sources(ws: Worksheet, result: DCFResult, generated_at: datetime) -> None:
    is_sample = result.historical == CONTROLLED_SAMPLE_HISTORICALS
    _title(ws, "Data Sources and Limitations",
           "Where the numbers in this workbook come from and what the model does not do.")
    _header(ws, 4, ["Topic", "Detail"], left=(0, 1))
    if is_sample:
        hist_status = (
            "CONTROLLED SAMPLE DATA. The historical figures are illustrative sample data defined "
            "in the DCF engine for demonstration and testing. They are not live market data and "
            "are not taken from any company filing."
        )
    else:
        hist_status = (
            "Custom historical data supplied by the caller of the DCF engine. It has not been "
            "verified by this model and is not live market data."
        )
    rows = [
        ("Historical data source / status", hist_status),
        ("Model assumptions source / status",
         "User-supplied inputs (share price, growth, margin, tax rate, WACC, terminal growth, net "
         "debt, shares outstanding). They are not derived from live market data; the current share "
         "price is an input, not a quote."),
        ("Live data / external services", "None. No market-data APIs or databases are used."),
        ("DCF methodology",
         "Unlevered DCF. Revenue grows at a constant rate; EBITDA = revenue x margin; EBIT = EBITDA - "
         "D&A; tax = EBIT x tax rate; NOPAT = EBIT - tax; UFCF = NOPAT + D&A - CapEx - change in NWC."),
        ("Discounting and terminal value",
         "End-of-year discounting: DF = 1 / (1 + WACC)^t. Terminal value (Gordon growth) = final-year "
         "UFCF x (1 + g) / (WACC - g), discounted with the final-year factor."),
        ("Valuation bridge",
         "Enterprise value = sum of PV of UFCF + PV of terminal value. Equity value = enterprise value - "
         "net debt. Implied share price = equity value / shares outstanding."),
        ("Source of truth",
         "Valuation results are calculated by the tested Python engine (backend/dcf.py). This workbook "
         "re-performs the calculation with live Excel formulas; 07_Model_Checks reconciles the two."),
        ("Generation timestamp", generated_at),
        ("Limitation: forecast simplification",
         "D&A, CapEx and change in NWC are held flat at their latest historical value; revenue growth, "
         "EBITDA margin and tax rate are constant across all forecast years."),
        ("Limitation: timing", "Cash flows are discounted at year end; no mid-year convention."),
        ("Limitation: tax", "Tax = EBIT x tax rate with no loss carry-forwards, so negative EBIT would "
                            "produce a tax credit."),
        ("Limitation: terminal value",
         "Single-stage perpetuity growth only; the terminal value is typically a large share of "
         "enterprise value (see the memo line on 04_DCF)."),
        ("Limitation: capital structure", "Net debt is a single static input; no leases, minorities, "
                                          "options or other adjustments."),
        ("Limitation: units", "Monetary amounts share one unspecified currency unit; shares outstanding "
                              "must be on the same scale."),
        ("Limitation: forecast length", "The number of forecast columns is fixed when the workbook is "
                                        "generated; editing the assumption does not add columns."),
        ("Status", "Illustrative automated DCF model. Not investment advice or a recommendation."),
    ]
    for i, (topic, detail) in enumerate(rows):
        r = 5 + i
        _put(ws, f"A{r}", topic, kind="label", bold=True, wrap=True)
        if isinstance(detail, datetime):
            _put(ws, f"B{r}", detail, kind="label", fmt="yyyy-mm-dd hh:mm:ss", align="left")
        else:
            _put(ws, f"B{r}", detail, kind="label", wrap=True)
    _widths(ws, {"A": 38, "B": 120})


# ---------------------------------------------------------------------------
# 00_Cover
# ---------------------------------------------------------------------------


def _build_cover(
    ws: Worksheet,
    a: dict[str, str],
    dcf: dict[str, str],
    generated_at: datetime,
) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = _NAVY
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    for col in range(1, 4):
        for row in (1, 2, 3):
            ws.cell(row=row, column=col).fill = PatternFill("solid", start_color=_NAVY, end_color=_NAVY)
    _put(ws, "A2", "DCF Valuation", kind="label", bold=True, size=20)
    ws["A2"].font = Font(name=_FONT, size=20, bold=True, color="FFFFFF")
    _put(ws, "A3", "Illustrative automated discounted cash flow model", kind="label")
    ws["A3"].font = Font(name=_FONT, size=10, italic=True, color="FFFFFF")

    rows = [
        ("Company", f"={a['company_name']}", None, False),
        ("Ticker", f"={a['ticker']}", None, False),
        ("Current share price", f"={a['share_price']}", FMT_PRICE, False),
        ("Implied share price", f"={dcf['price']}", FMT_PRICE, True),
        ("Upside / (downside)", f"={dcf['upside']}", FMT_PCT1, True),
        ("Enterprise value", f"={dcf['ev']}", FMT_NUM, True),
        ("Equity value", f"={dcf['equity']}", FMT_NUM, True),
        ("Forecast period", f"={a['forecast_years']}", FMT_YEARS, False),
        ("Model date", generated_at, "yyyy-mm-dd hh:mm", False),
        ("Model checks", f"={_abs(SHEET_CHECKS, 'B4')}", None, False),
    ]
    for i, (label, value, fmt, key) in enumerate(rows):
        r = 5 + i
        _put(ws, f"A{r}", label, kind="label", bold=key)
        kind = "label" if label == "Model date" else "link"
        _put(ws, f"B{r}", value, kind=kind, fmt=fmt, bold=key, fill=_KEY_FILL if key else None,
             border=_BOX, align="right")
    notes = [
        "This is an illustrative automated DCF model generated from controlled sample data.",
        "It is not investment advice, a price target or a recommendation.",
        "Amounts are in a single unspecified currency unit. See 08_Data_Sources for sources and limitations.",
        "Sheets: 01_Assumptions, 02_Historical, 03_Forecast, 04_DCF, 05_Sensitivity, "
        "06_Valuation_Summary, 07_Model_Checks, 08_Data_Sources.",
    ]
    for i, note in enumerate(notes):
        _put(ws, f"A{17 + i}", note, kind="label", italic=True)
    _widths(ws, {"A": 26, "B": 32, "C": 60})


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def build_dcf_workbook(result: DCFResult, generated_at: datetime | None = None) -> Workbook:
    """Build the workbook in memory from an already-calculated ``DCFResult``."""
    generated_at = generated_at or datetime.now().replace(microsecond=0)

    wb = Workbook()
    wb.active.title = SHEET_NAMES[0]
    for name in SHEET_NAMES[1:]:
        wb.create_sheet(name)
    ws = {name: wb[name] for name in SHEET_NAMES}

    assumptions = _build_assumptions(ws[SHEET_ASSUMPTIONS], result)
    historical = _build_historical(ws[SHEET_HISTORICAL], result)
    forecast = _build_forecast(ws[SHEET_FORECAST], result, assumptions, historical)
    dcf = _build_dcf(ws[SHEET_DCF], result, assumptions, forecast)
    _build_sensitivity(ws[SHEET_SENSITIVITY], result, assumptions, dcf)
    _build_summary(ws[SHEET_SUMMARY], assumptions, dcf)
    _build_checks(ws[SHEET_CHECKS], result, assumptions, dcf, forecast)
    _build_sources(ws[SHEET_SOURCES], result, generated_at)
    _build_cover(ws[SHEET_COVER], assumptions, dcf, generated_at)

    wb.properties.title = f"DCF Valuation - {result.inputs.company_name}"
    wb.properties.creator = "Financial Analyst Automation Platform"
    wb.calculation = CalcProperties(fullCalcOnLoad=True)  # calculate formulas on open
    return wb


def _resolve_output_path(
    output: str | Path | None, ticker: str, generated_at: datetime
) -> Path:
    """Turn an output file path or output directory into a concrete .xlsx path."""
    target = Path(output) if output is not None else DEFAULT_OUTPUT_DIR
    if not target.is_dir() and target.suffix:
        if target.suffix.lower() != ".xlsx":
            raise ValueError("Output path must end with .xlsx or be a directory.")
        path = target
    else:
        safe_ticker = re.sub(r"[^A-Za-z0-9]+", "", ticker) or "DCF"
        path = target / f"{safe_ticker}_DCF_{generated_at:%Y%m%d_%H%M%S}.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def generate_dcf_workbook(
    result: DCFResult,
    output: str | Path | None = None,
    generated_at: datetime | None = None,
) -> Path:
    """Generate the DCF workbook and save it; return the path of the saved file.

    Args:
        result: Output of ``run_dcf`` (the single source of truth).
        output: A ``.xlsx`` file path, or a directory in which a file named
            ``<TICKER>_DCF_<timestamp>.xlsx`` is created. Defaults to the
            project's ``outputs`` folder.
        generated_at: Model date shown in the workbook (defaults to now).
    """
    generated_at = generated_at or datetime.now().replace(microsecond=0)
    path = _resolve_output_path(output, result.inputs.ticker, generated_at)
    build_dcf_workbook(result, generated_at).save(path)
    return path