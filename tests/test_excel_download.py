"""Stage 5D tests for browser-triggered Excel model generation/download."""

from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import backend.main as main_module
from backend.main import app

client = TestClient(app)


def valid_payload(**overrides):
    payload = {
        "company_name": "Apple Inc.",
        "ticker": "AAPL",
        "share_price": 255,
        "forecast_years": 5,
        "revenue_growth": 0.08,
        "ebitda_margin": 0.33,
        "tax_rate": 0.16,
        "wacc": 0.085,
        "terminal_growth": 0.025,
        "net_debt": 50000,
        "shares_outstanding": 15000,
    }
    payload.update(overrides)
    return payload


@pytest.fixture
def isolated_outputs(tmp_path, monkeypatch):
    monkeypatch.setattr(main_module, "OUTPUTS_DIR", tmp_path)
    return tmp_path


def test_excel_endpoint_returns_xlsx(isolated_outputs):
    response = client.post("/api/dcf/excel", json=valid_payload())
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.content[:2] == b"PK"


def test_excel_download_has_professional_filename(isolated_outputs):
    response = client.post("/api/dcf/excel", json=valid_payload())
    disposition = response.headers["content-disposition"]
    assert "attachment" in disposition.lower()
    assert "AAPL_DCF_" in disposition
    assert ".xlsx" in disposition


def test_excel_endpoint_creates_one_file_in_outputs(isolated_outputs):
    response = client.post("/api/dcf/excel", json=valid_payload())
    assert response.status_code == 200
    files = list(isolated_outputs.glob("*.xlsx"))
    assert len(files) == 1
    assert files[0].name.startswith("AAPL_DCF_")


def test_downloaded_workbook_has_required_nine_sheets(isolated_outputs):
    response = client.post("/api/dcf/excel", json=valid_payload())
    wb = load_workbook(BytesIO(response.content), data_only=False, read_only=True)
    assert wb.sheetnames == [
        "00_Cover",
        "01_Assumptions",
        "02_Historical",
        "03_Forecast",
        "04_DCF",
        "05_Sensitivity",
        "06_Valuation_Summary",
        "07_Model_Checks",
        "08_Data_Sources",
    ]
    wb.close()


def test_downloaded_workbook_contains_submitted_company_and_ticker(isolated_outputs):
    response = client.post("/api/dcf/excel", json=valid_payload(company_name="Demo Corp", ticker="DEMO"))
    wb = load_workbook(BytesIO(response.content), data_only=False, read_only=True)
    assumptions = wb["01_Assumptions"]
    values = [cell.value for row in assumptions.iter_rows() for cell in row]
    assert "Demo Corp" in values
    assert "DEMO" in values
    wb.close()


def test_excel_endpoint_preserves_financial_validation(isolated_outputs):
    response = client.post(
        "/api/dcf/excel",
        json=valid_payload(wacc=0.02, terminal_growth=0.025),
    )
    assert response.status_code == 422
    assert "WACC must be greater than terminal growth" in response.text
    assert list(isolated_outputs.glob("*.xlsx")) == []


def test_frontend_exposes_excel_download_workflow():
    root = client.get("/")
    javascript = client.get("/static/app.js")
    assert root.status_code == 200
    assert 'id="download-excel"' in root.text
    assert "Download Excel Model" in root.text
    assert javascript.status_code == 200
    assert 'fetch("/api/dcf/excel"' in javascript.text
    assert "Generating Excel..." in javascript.text
