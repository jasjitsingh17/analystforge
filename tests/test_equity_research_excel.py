"""Stage 6D-A/B tests for the Equity Research Excel workbook generator."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

from backend.equity_research import EquityResearchInputs, run_equity_research
from backend.equity_research_excel import (
    CHECK_TOLERANCE,
    DEFAULT_OUTPUT_DIR,
    SHEET_CHECKS,
    SHEET_CONSENSUS,
    SHEET_COVER,
    SHEET_GUIDANCE,
    SHEET_HISTORICAL,
    SHEET_KPIS,
    SHEET_NAMES,
    SHEET_SOURCES,
    SHEET_SUMMARY,
    SHEET_TAKEAWAYS,
    SHEET_VALUATION,
    build_equity_research_workbook,
    generate_equity_research_workbook,
)

GENERATED_AT = datetime(2026, 9, 25, 17, 0, 0)


@pytest.fixture
def result():
    return run_equity_research(EquityResearchInputs())


@pytest.fixture
def workbook_path(tmp_path, result):
    return generate_equity_research_workbook(result, tmp_path / "NST_ER_Model.xlsx", GENERATED_AT)


@pytest.fixture
def wb(workbook_path):
    return load_workbook(workbook_path, data_only=False)


# File / workbook structure ----------------------------------------------------------

def test_workbook_is_created(workbook_path):
    assert workbook_path.exists()


def test_output_is_real_xlsx(workbook_path):
    assert workbook_path.suffix == ".xlsx"
    with ZipFile(workbook_path) as archive:
        assert "xl/workbook.xml" in archive.namelist()


def test_workbook_opens(workbook_path):
    loaded = load_workbook(workbook_path, data_only=False)
    assert loaded.sheetnames


def test_exact_sheet_names(wb):
    assert tuple(wb.sheetnames) == SHEET_NAMES


def test_exact_sheet_count(wb):
    assert len(wb.sheetnames) == 10


@pytest.mark.parametrize("sheet_name", SHEET_NAMES)
def test_gridlines_are_hidden(wb, sheet_name):
    assert wb[sheet_name].sheet_view.showGridLines is False


def test_workbook_requests_full_recalculation(result):
    workbook = build_equity_research_workbook(result, GENERATED_AT)
    assert workbook.calculation.fullCalcOnLoad is True
    assert workbook.calculation.forceFullCalc is True
    assert workbook.calculation.calcMode == "auto"


# Cover ------------------------------------------------------------------------------

def test_cover_is_formula_driven(wb):
    ws = wb[SHEET_COVER]
    assert ws["B6"].value == "='01_Earnings_Summary'!B5"
    assert ws["B9"].value == "='01_Earnings_Summary'!B19"
    assert ws["B18"].value == "='06_Valuation'!B16"
    assert ws["B24"].value == "='08_Model_Checks'!B4"


def test_cover_discloses_illustrative_nature(wb):
    ws = wb[SHEET_COVER]
    text = " ".join(str(ws[cell].value or "") for cell in ["A4", "A27", "A29"])
    lower = text.lower()
    assert "illustrative" in lower
    assert "controlled sample" in lower
    assert "not investment advice" in lower


def test_cover_records_generation_timestamp(wb):
    assert wb[SHEET_COVER]["B31"].value == GENERATED_AT


# Earnings summary inputs and formulas ----------------------------------------------

@pytest.mark.parametrize(
    "cell, expected",
    [
        ("B5", "Northstar Technologies plc"),
        ("B6", "NST"),
        ("B7", "GBP"),
        ("B8", "FY2026 H1"),
    ],
)
def test_summary_identity_fields(wb, cell, expected):
    assert wb[SHEET_SUMMARY][cell].value == expected


@pytest.mark.parametrize(
    "cell, expected",
    [
        ("B12", 1100.0), ("C12", 1200.0), ("D12", 1260.0),
        ("B13", 240.0), ("C13", 280.0), ("D13", 300.0),
        ("B14", 0.36), ("C14", 0.42), ("D14", 0.46),
        ("B15", 170.0), ("C15", 195.0), ("D15", 210.0),
    ],
)
def test_summary_raw_earnings_inputs_match_engine(wb, cell, expected):
    assert wb[SHEET_SUMMARY][cell].value == expected


@pytest.mark.parametrize("row", [12, 13, 14, 15])
def test_summary_yoy_cells_are_formulas(wb, row):
    assert wb[SHEET_SUMMARY][f"E{row}"].value == f"=(D{row}-B{row})/ABS(B{row})"


@pytest.mark.parametrize("row", [12, 13, 14, 15])
def test_summary_surprise_cells_are_formulas(wb, row):
    assert wb[SHEET_SUMMARY][f"F{row}"].value == f"=(D{row}-C{row})/ABS(C{row})"


@pytest.mark.parametrize("row", [12, 13, 14, 15])
def test_summary_classification_cells_are_formulas(wb, row):
    formula = wb[SHEET_SUMMARY][f"G{row}"].value
    assert formula.startswith("=IF(")
    assert "BEAT" in formula and "MISS" in formula and "IN-LINE" in formula


@pytest.mark.parametrize("cell", ["B5", "B12", "C12", "D12", "B13", "C13", "D13"])
def test_hardcoded_summary_inputs_are_blue(wb, cell):
    assert wb[SHEET_SUMMARY][cell].font.color.rgb[-6:] == "0000FF"


@pytest.mark.parametrize("cell", ["B12", "C12", "D12", "B13", "C13", "D13"])
def test_raw_inputs_have_source_comments(wb, cell):
    comment = wb[SHEET_SUMMARY][cell].comment
    assert comment is not None
    assert "Controlled sample data" in comment.text


# Historical / expectation sheets ----------------------------------------------------

def test_historical_sheet_links_to_summary(wb):
    ws = wb[SHEET_HISTORICAL]
    assert ws["B6"].value == "='01_Earnings_Summary'!B12"
    assert ws["C6"].value == "='01_Earnings_Summary'!D12"
    assert ws["B8"].value == "='01_Earnings_Summary'!B13/'01_Earnings_Summary'!B12"
    assert ws["C11"].value == "='01_Earnings_Summary'!D15/'01_Earnings_Summary'!D13"


def test_consensus_sheet_links_and_calculates_variances(wb):
    ws = wb[SHEET_CONSENSUS]
    assert ws["B6"].value == "='01_Earnings_Summary'!C12"
    assert ws["C6"].value == "='01_Earnings_Summary'!D12"
    assert ws["D6"].value == "=C6-B6"
    assert ws["E6"].value == "=(C6-B6)/ABS(B6)"
    assert ws["F6"].value == "='01_Earnings_Summary'!G12"


def test_consensus_sheet_contains_chart(wb):
    assert len(wb[SHEET_CONSENSUS]._charts) >= 1


# KPIs -------------------------------------------------------------------------------

@pytest.mark.parametrize(
    "cell, fragment",
    [
        ("B6", "E12"),
        ("B7", "E13"),
        ("B8", "E14"),
        ("B9", "E15"),
        ("B15", "C8"),
        ("B18", "D15"),
    ],
)
def test_kpi_sheet_is_formula_driven(wb, cell, fragment):
    value = wb[SHEET_KPIS][cell].value
    assert isinstance(value, str) and value.startswith("=")
    assert fragment in value


def test_margin_surprise_is_calculated_in_bps(wb):
    assert wb[SHEET_KPIS]["B16"].value == "=(B15-B14)*10000"
    assert "bps" in wb[SHEET_KPIS]["B16"].number_format


# Guidance ---------------------------------------------------------------------------

@pytest.mark.parametrize(
    "cell, expected",
    [
        ("B6", 4900.0), ("C6", 5100.0), ("E6", 5000.0), ("F6", 5200.0),
        ("B7", 1000.0), ("C7", 1080.0), ("E7", 1050.0), ("F7", 1130.0),
    ],
)
def test_guidance_raw_inputs_match_engine(wb, cell, expected):
    assert wb[SHEET_GUIDANCE][cell].value == expected


@pytest.mark.parametrize("row", [6, 7])
def test_guidance_midpoints_and_revisions_are_formulas(wb, row):
    ws = wb[SHEET_GUIDANCE]
    assert ws[f"D{row}"].value == f"=AVERAGE(B{row}:C{row})"
    assert ws[f"G{row}"].value == f"=AVERAGE(E{row}:F{row})"
    assert ws[f"H{row}"].value == f"=(G{row}-D{row})/D{row}"


def test_guidance_classification_formulas_exist(wb):
    assert "RAISED" in wb[SHEET_GUIDANCE]["B11"].value
    assert "LOWERED" in wb[SHEET_GUIDANCE]["B12"].value


def test_guidance_sheet_contains_chart(wb):
    assert len(wb[SHEET_GUIDANCE]._charts) >= 1


# Valuation --------------------------------------------------------------------------

@pytest.mark.parametrize(
    "cell, expected",
    [
        ("B5", 25.0),
        ("B6", 200.0),
        ("B7", 600.0),
        ("B8", 5200.0),
        ("B9", 1100.0),
        ("B10", 1.60),
    ],
)
def test_valuation_inputs_match_engine(wb, cell, expected):
    assert wb[SHEET_VALUATION][cell].value == expected


@pytest.mark.parametrize(
    "cell, formula",
    [
        ("B14", "=B5*B6"),
        ("B15", "=B14+B7"),
        ("B16", "=B5/B10"),
        ("B17", "=B15/B9"),
        ("B18", "=B15/B8"),
        ("B19", "=B7/B9"),
    ],
)
def test_valuation_outputs_are_formulas(wb, cell, formula):
    assert wb[SHEET_VALUATION][cell].value == formula


def test_valuation_sheet_contains_chart(wb):
    assert len(wb[SHEET_VALUATION]._charts) >= 1


def test_multiple_cells_use_x_number_format(wb):
    for cell in ["B16", "B17", "B18", "B19"]:
        assert "x" in wb[SHEET_VALUATION][cell].number_format


# Analyst commentary -----------------------------------------------------------------

def test_takeaways_match_python_engine(wb, result):
    ws = wb[SHEET_TAKEAWAYS]
    workbook_text = " ".join(str(ws.cell(row=r, column=2).value or "") for r in range(5, 8))
    for takeaway in result.takeaways:
        assert takeaway in workbook_text


def test_catalysts_match_python_engine(wb, result):
    text = " ".join(str(cell.value or "") for row in wb[SHEET_TAKEAWAYS].iter_rows() for cell in row)
    for catalyst in result.catalysts:
        assert catalyst in text


def test_risks_match_python_engine(wb, result):
    text = " ".join(str(cell.value or "") for row in wb[SHEET_TAKEAWAYS].iter_rows() for cell in row)
    for risk in result.risks:
        assert risk in text


# Model checks ------------------------------------------------------------------------

CHECK_EXPECTATIONS = [
    (8, "Revenue surprise", "revenue", "surprise"),
    (9, "EBITDA surprise", "ebitda", "surprise"),
    (10, "EPS surprise", "eps", "surprise"),
    (11, "FCF surprise", "fcf", "surprise"),
]


@pytest.mark.parametrize("row,label,metric_name,attribute", CHECK_EXPECTATIONS)
def test_model_checks_hold_python_metric_references(wb, result, row, label, metric_name, attribute):
    ws = wb[SHEET_CHECKS]
    assert ws[f"B{row}"].value == label
    expected = getattr(getattr(result, metric_name), attribute)
    assert ws[f"C{row}"].value == pytest.approx(expected)


@pytest.mark.parametrize(
    "row, expected",
    [
        (12, 300 / 1260),
        (13, ((300 / 1260) - (280 / 1200)) * 10000),
        (14, 5000.0),
        (15, 5600.0),
        (16, 15.625),
        (17, 5600 / 1100),
        (18, 0.02),
        (19, (1090 - 1040) / 1040),
    ],
)
def test_remaining_python_check_values_are_exact(wb, row, expected):
    assert wb[SHEET_CHECKS][f"C{row}"].value == pytest.approx(expected)


@pytest.mark.parametrize("row", range(8, 20))
def test_each_model_check_has_excel_link_difference_and_status_formula(wb, row):
    ws = wb[SHEET_CHECKS]
    assert isinstance(ws[f"D{row}"].value, str) and ws[f"D{row}"].value.startswith("=")
    assert ws[f"E{row}"].value == f"=D{row}-C{row}"
    assert ws[f"F{row}"].value == f'=IF(ABS(E{row})<=$B$5,"PASS","FAIL")'


def test_overall_model_status_formula(wb):
    assert wb[SHEET_CHECKS]["B4"].value == '=IF(COUNTIF(F8:F19,"FAIL")=0,"ALL CHECKS PASS","CHECK MODEL")'


def test_check_tolerance_is_tight(wb):
    assert wb[SHEET_CHECKS]["B5"].value == CHECK_TOLERANCE
    assert CHECK_TOLERANCE <= 1e-6


def test_python_reference_values_are_purple(wb):
    for row in range(8, 20):
        assert wb[SHEET_CHECKS][f"C{row}"].font.color.rgb[-6:] == "7030A0"


def test_cross_sheet_excel_values_are_green(wb):
    for row in range(8, 20):
        assert wb[SHEET_CHECKS][f"D{row}"].font.color.rgb[-6:] == "008000"


# Data sources / disclosures ----------------------------------------------------------

def test_data_sources_discloses_controlled_sample_data(wb):
    text = " ".join(str(cell.value or "") for row in wb[SHEET_SOURCES].iter_rows() for cell in row).lower()
    assert "controlled sample" in text
    assert "fictional issuer" in text
    assert "not live" in text or "not company-reported live data" in text


def test_data_sources_disclaims_investment_advice(wb):
    text = " ".join(str(cell.value or "") for row in wb[SHEET_SOURCES].iter_rows() for cell in row).lower()
    assert "not investment advice" in text
    assert "not" in text and "price target" in text


def test_data_sources_records_generation_timestamp(wb):
    assert wb[SHEET_SOURCES]["B23"].value == GENERATED_AT


# Formatting conventions -------------------------------------------------------------

def test_zero_and_negative_finance_number_format_convention(wb):
    fmt = wb[SHEET_VALUATION]["B14"].number_format
    assert "[Red]" in fmt
    assert "-" in fmt


def test_cross_sheet_links_are_green(wb):
    assert wb[SHEET_HISTORICAL]["B6"].font.color.rgb[-6:] == "008000"


def test_formula_cells_are_black(wb):
    assert wb[SHEET_VALUATION]["B14"].font.color.rgb[-6:] == "000000"


# Output path handling ---------------------------------------------------------------

def test_directory_output_creates_professional_filename(tmp_path, result):
    path = generate_equity_research_workbook(result, tmp_path, GENERATED_AT)
    assert path.parent == tmp_path
    assert path.name == "NST_ER_FY2026_H1_20260925_170000.xlsx"
    assert path.exists()


def test_explicit_xlsx_output_path_is_respected(tmp_path, result):
    target = tmp_path / "custom" / "research.xlsx"
    path = generate_equity_research_workbook(result, target, GENERATED_AT)
    assert path == target
    assert target.exists()


def test_missing_output_directories_are_created(tmp_path, result):
    target = tmp_path / "a" / "b" / "model.xlsx"
    path = generate_equity_research_workbook(result, target, GENERATED_AT)
    assert path.exists()


def test_non_xlsx_output_is_rejected(tmp_path, result):
    with pytest.raises(ValueError, match=r"\.xlsx"):
        generate_equity_research_workbook(result, tmp_path / "model.csv", GENERATED_AT)


def test_default_output_directory_is_project_outputs_folder():
    project_root = Path(__file__).resolve().parent.parent
    # In the user's project this module sits under backend/ and resolves to project/outputs.
    assert DEFAULT_OUTPUT_DIR.name == "outputs"


# Alternate inputs / adaptability ----------------------------------------------------

def test_generator_adapts_to_changed_company_and_ticker(tmp_path):
    inputs = EquityResearchInputs(company_name="Atlas Systems plc", ticker="ATS", period="FY2027")
    result = run_equity_research(inputs)
    path = generate_equity_research_workbook(result, tmp_path, GENERATED_AT)
    assert path.name.startswith("ATS_ER_FY2027_")
    loaded = load_workbook(path, data_only=False)
    assert loaded[SHEET_SUMMARY]["B5"].value == "Atlas Systems plc"
    assert loaded[SHEET_SUMMARY]["B6"].value == "ATS"


def test_generator_adapts_to_miss_case(tmp_path):
    inputs = EquityResearchInputs(actual_revenue=1150.0)
    result = run_equity_research(inputs)
    path = generate_equity_research_workbook(result, tmp_path / "miss.xlsx", GENERATED_AT)
    loaded = load_workbook(path, data_only=False)
    # Excel formula remains formula-driven; Python reference on checks reflects changed engine.
    assert loaded[SHEET_SUMMARY]["D12"].value == 1150.0
    assert loaded[SHEET_CHECKS]["C8"].value == pytest.approx(result.revenue.surprise)
    assert result.revenue.classification == "MISS"


def test_generator_adapts_to_net_cash_case(tmp_path):
    inputs = EquityResearchInputs(net_debt=-400.0)
    result = run_equity_research(inputs)
    path = generate_equity_research_workbook(result, tmp_path / "net_cash.xlsx", GENERATED_AT)
    loaded = load_workbook(path, data_only=False)
    assert loaded[SHEET_VALUATION]["B7"].value == -400.0
    assert loaded[SHEET_CHECKS]["C15"].value == pytest.approx(result.enterprise_value)
