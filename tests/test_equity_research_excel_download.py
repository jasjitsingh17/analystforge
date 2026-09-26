"""Stage 6D-C tests for the Equity Research Excel download workflow."""

from dataclasses import asdict
from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

import backend.main as main_module
from backend.equity_research import CONTROLLED_SAMPLE_INPUTS
from backend.main import app

client = TestClient(app)
ER_EXCEL_URL = "/api/equity-research/excel"
EXPECTED_SHEETS = [
    "00_Cover",
    "01_Earnings_Summary",
    "02_Historical",
    "03_Consensus_vs_Actual",
    "04_KPIs",
    "05_Guidance",
    "06_Valuation",
    "07_Analyst_Takeaways",
    "08_Model_Checks",
    "09_Data_Sources",
]


def payload(**overrides):
    data = asdict(CONTROLLED_SAMPLE_INPUTS)
    data.update(overrides)
    return data


def test_er_excel_endpoint_returns_real_xlsx(tmp_path, monkeypatch):
    monkeypatch.setattr(main_module, "OUTPUTS_DIR", tmp_path)
    response = client.post(ER_EXCEL_URL, json=payload())
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.content[:2] == b"PK"


def test_er_excel_download_has_professional_filename(tmp_path, monkeypatch):
    monkeypatch.setattr(main_module, "OUTPUTS_DIR", tmp_path)
    response = client.post(ER_EXCEL_URL, json=payload())
    disposition = response.headers.get("content-disposition", "")
    assert "attachment" in disposition.lower()
    assert "NST_ER_FY2026_H1_" in disposition
    assert ".xlsx" in disposition


def test_er_excel_endpoint_creates_one_file_in_outputs(tmp_path, monkeypatch):
    monkeypatch.setattr(main_module, "OUTPUTS_DIR", tmp_path)
    response = client.post(ER_EXCEL_URL, json=payload())
    assert response.status_code == 200
    files = list(tmp_path.glob("*.xlsx"))
    assert len(files) == 1
    assert files[0].stat().st_size > 0


def test_downloaded_er_workbook_has_exact_ten_sheet_structure(tmp_path, monkeypatch):
    monkeypatch.setattr(main_module, "OUTPUTS_DIR", tmp_path)
    response = client.post(ER_EXCEL_URL, json=payload())
    wb = load_workbook(BytesIO(response.content), data_only=False)
    assert wb.sheetnames == EXPECTED_SHEETS


def test_downloaded_er_workbook_contains_submitted_company_ticker_and_period(tmp_path, monkeypatch):
    monkeypatch.setattr(main_module, "OUTPUTS_DIR", tmp_path)
    response = client.post(
        ER_EXCEL_URL,
        json=payload(company_name="Example Holdings plc", ticker="EXH", period="Q3 2026"),
    )
    wb = load_workbook(BytesIO(response.content), data_only=False)
    summary = wb["01_Earnings_Summary"]
    values = [cell.value for row in summary.iter_rows() for cell in row]
    assert "Example Holdings plc" in values
    assert "EXH" in values
    assert "Q3 2026" in values


def test_er_excel_endpoint_preserves_financial_validation():
    response = client.post(ER_EXCEL_URL, json=payload(share_price=0))
    assert response.status_code == 422
    assert "Share price" in response.json()["detail"][0]["msg"]


def test_er_excel_route_is_exposed_in_openapi():
    assert ER_EXCEL_URL in client.get("/openapi.json").json()["paths"]


def test_frontend_exposes_research_model_download_button():
    html = client.get("/").text
    assert 'id="download-er-excel"' in html
    assert "Download Research Model" in html
    assert "disabled" in html


def test_frontend_javascript_uses_er_excel_endpoint_and_stale_input_guard():
    js = client.get("/static/app.js").text
    assert 'fetch("/api/equity-research/excel"' in js
    assert "lastSuccessfulERPayload" in js
    assert "Inputs changed" in js
    assert "Generating Research Model..." in js
