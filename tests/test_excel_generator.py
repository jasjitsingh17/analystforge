"""Tests for backend/excel_generator.py.

How the Excel numbers are verified
----------------------------------
openpyxl writes formulas but cannot calculate them, so a workbook it just
wrote has no cached values. To test what Excel would actually calculate, this
file contains ``FormulaEvaluator``: a small evaluator for the limited formula
grammar the generator uses (cell and range references, + - * / ^, comparisons,
SUM, COUNT, COUNTIF, IF, AND, ABS, ROUND). It reads the formulas stored in the
generated workbook and evaluates them cell by cell.

The reference for every reconciliation is ``run_dcf()`` called inside this
test file. The workbook's own "Python" column on 07_Model_Checks is never
used as the expected value, so nothing is compared with a copy of itself.

Additional evidence that is not part of this suite: the same workbook was
recalculated with LibreOffice during development and matched the Python
engine (see the Stage 4A report).
"""

from __future__ import annotations

import math
import re
from dataclasses import replace
from datetime import datetime
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from backend.dcf import DCFInputs, run_dcf
from backend.excel_generator import DEFAULT_OUTPUT_DIR, generate_dcf_workbook

GENERATED_AT = datetime(2026, 9, 24, 10, 30, 0)

REQUIRED_SHEETS = [
    "00_Cover",
    "01_Assumptions",
    "02_Historical",
    "03_Forecast",
    "04_DCF",
    "05_Sensitivity",
    "06_Valuation_Summary",
    "07_Model_Checks",
    "08_Data_Sources",
]

CHECK_NAMES = [
    "Python vs Excel: Enterprise value",
    "Python vs Excel: Equity value",
    "Python vs Excel: Implied share price",
    "Forecast year count",
    "Enterprise value bridge",
    "Equity value bridge",
    "Implied share price reconciliation",
    "WACC greater than terminal growth",
    "Python vs Excel: Sum of PV of UFCF",
    "Python vs Excel: Terminal value",
    "Python vs Excel: PV of terminal value",
    "Python vs Excel: Total UFCF",
    "Forecast years assumption vs forecast columns",
]

# Sensitivity axes required by the specification (base case in the centre).
WACC_AXIS = [0.075, 0.08, 0.085, 0.09, 0.095]
GROWTH_AXIS = [0.015, 0.02, 0.025, 0.03, 0.035]
GRID_COLUMNS = ["C", "D", "E", "F", "G"]
GRID_ROWS = [8, 9, 10, 11, 12]


# ---------------------------------------------------------------------------
# Tolerance helper (as requested: rel_tol=1e-9, abs_tol=1e-6)
# ---------------------------------------------------------------------------


def close(actual: Any, expected: Any, rel_tol: float = 1e-9, abs_tol: float = 1e-6) -> bool:
    if isinstance(actual, bool) or isinstance(expected, bool):
        return False
    if not isinstance(actual, (int, float)) or not isinstance(expected, (int, float)):
        return False
    return math.isclose(actual, expected, rel_tol=rel_tol, abs_tol=abs_tol)


def assert_close(actual: Any, expected: Any, what: str = "value") -> None:
    assert close(actual, expected), f"{what}: workbook={actual!r} expected={expected!r}"


def assert_all_close(actual: list[Any], expected: list[Any], what: str = "series") -> None:
    assert len(actual) == len(expected), f"{what}: length {len(actual)} != {len(expected)}"
    for i, (a, e) in enumerate(zip(actual, expected), start=1):
        assert_close(a, e, f"{what}[{i}]")


# ---------------------------------------------------------------------------
# Formula evaluator (test infrastructure)
# ---------------------------------------------------------------------------

_TOKEN = re.compile(
    r"""\s*(?:
        (?P<string>"(?:[^"]|"")*")
      | (?P<ref>(?:'[^']+'!|[A-Za-z0-9_]+!)?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)
      | (?P<func>[A-Z][A-Z0-9.]*(?=\())
      | (?P<number>\d+\.?\d*(?:[eE][+-]?\d+)?|\.\d+)
      | (?P<op><=|>=|<>|[-+*/^=<>(),])
    )""",
    re.VERBOSE,
)


def _tokenize(text: str) -> list[tuple[str, str]]:
    tokens: list[tuple[str, str]] = []
    pos = 0
    text = text.strip()
    while pos < len(text):
        match = _TOKEN.match(text, pos)
        if not match or match.end() == pos:
            raise ValueError(f"Cannot parse formula near {text[pos:]!r}")
        pos = match.end()
        kind = match.lastgroup
        tokens.append((kind, match.group(kind)))
    return tokens


class _Parser:
    """Recursive-descent parser using Excel precedence:
    comparison < + - < * / < ^ < unary minus (so -2^2 = 4, as in Excel)."""

    _COMPARISONS = {"=", "<>", "<", ">", "<=", ">="}

    def __init__(self, tokens: list[tuple[str, str]]) -> None:
        self.tokens = tokens
        self.pos = 0

    def _peek(self) -> tuple[str | None, str | None]:
        return self.tokens[self.pos] if self.pos < len(self.tokens) else (None, None)

    def _take(self) -> tuple[str | None, str | None]:
        token = self._peek()
        self.pos += 1
        return token

    def parse(self) -> tuple:
        node = self._comparison()
        if self.pos != len(self.tokens):
            raise ValueError("Unexpected trailing tokens in formula")
        return node

    def _comparison(self) -> tuple:
        left = self._additive()
        while self._peek()[0] == "op" and self._peek()[1] in self._COMPARISONS:
            op = self._take()[1]
            left = ("cmp", op, left, self._additive())
        return left

    def _additive(self) -> tuple:
        left = self._term()
        while self._peek()[0] == "op" and self._peek()[1] in ("+", "-"):
            op = self._take()[1]
            left = ("bin", op, left, self._term())
        return left

    def _term(self) -> tuple:
        left = self._power()
        while self._peek()[0] == "op" and self._peek()[1] in ("*", "/"):
            op = self._take()[1]
            left = ("bin", op, left, self._power())
        return left

    def _power(self) -> tuple:
        left = self._unary()
        while self._peek() == ("op", "^"):
            self._take()
            left = ("bin", "^", left, self._unary())
        return left

    def _unary(self) -> tuple:
        if self._peek()[0] == "op" and self._peek()[1] in ("-", "+"):
            op = self._take()[1]
            operand = self._unary()
            return ("neg", operand) if op == "-" else operand
        return self._primary()

    def _primary(self) -> tuple:
        kind, value = self._take()
        if kind == "number":
            return ("num", float(value))
        if kind == "string":
            return ("str", value[1:-1].replace('""', '"'))
        if kind == "ref":
            return ("ref", value)
        if kind == "func":
            assert self._take() == ("op", "(")
            args: list[tuple] = []
            if self._peek() != ("op", ")"):
                while True:
                    args.append(self._comparison())
                    if self._peek() == ("op", ","):
                        self._take()
                    else:
                        break
            assert self._take() == ("op", ")"), "Missing closing parenthesis"
            return ("func", value, args)
        if (kind, value) == ("op", "("):
            node = self._comparison()
            assert self._take() == ("op", ")"), "Missing closing parenthesis"
            return node
        raise ValueError(f"Unexpected token {value!r}")


def _to_number(value: Any) -> float:
    if value is None:
        return 0.0
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    raise TypeError(f"#VALUE!: cannot use {value!r} in arithmetic")


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


class FormulaEvaluator:
    """Evaluates the formulas of an openpyxl workbook loaded WITH formulas."""

    def __init__(self, workbook: Any) -> None:
        self.wb = workbook
        self._cache: dict[tuple[str, str], Any] = {}
        self._active: set[tuple[str, str]] = set()

    def value(self, sheet: str, coordinate: str) -> Any:
        coordinate = coordinate.replace("$", "")
        key = (sheet, coordinate)
        if key in self._cache:
            return self._cache[key]
        raw = self.wb[sheet][coordinate].value
        if isinstance(raw, str) and raw.startswith("="):
            if key in self._active:
                raise RuntimeError(f"Circular reference at {sheet}!{coordinate}")
            self._active.add(key)
            try:
                result = self._eval(_Parser(_tokenize(raw[1:])).parse(), sheet)
            finally:
                self._active.discard(key)
        else:
            result = raw
        self._cache[key] = result
        return result

    def _eval(self, node: tuple, sheet: str) -> Any:
        tag = node[0]
        if tag in ("num", "str"):
            return node[1]
        if tag == "ref":
            return self._ref(node[1], sheet)
        if tag == "neg":
            return -_to_number(self._eval(node[1], sheet))
        if tag == "bin":
            _, op, left, right = node
            a = _to_number(self._eval(left, sheet))
            b = _to_number(self._eval(right, sheet))
            if op == "+":
                return a + b
            if op == "-":
                return a - b
            if op == "*":
                return a * b
            if op == "/":
                if b == 0:
                    raise ZeroDivisionError("#DIV/0!")
                return a / b
            return a ** b
        if tag == "cmp":
            _, op, left, right = node
            a = self._eval(left, sheet)
            b = self._eval(right, sheet)
            if isinstance(a, str) and isinstance(b, str):
                a, b = a.lower(), b.lower()
            elif isinstance(a, str) or isinstance(b, str):
                return op == "<>"  # text never equals a number
            else:
                a, b = _to_number(a), _to_number(b)
            return {"=": a == b, "<>": a != b, "<": a < b, ">": a > b,
                    "<=": a <= b, ">=": a >= b}[op]
        if tag == "func":
            return self._call(node[1], node[2], sheet)
        raise ValueError(f"Unknown node {tag}")

    def _ref(self, text: str, sheet: str) -> Any:
        if "!" in text:
            sheet_part, cell_part = text.rsplit("!", 1)
            sheet = sheet_part.strip("'")
        else:
            cell_part = text
        cell_part = cell_part.replace("$", "")
        if ":" in cell_part:
            min_col, min_row, max_col, max_row = range_boundaries(cell_part)
            return [
                self.value(sheet, f"{get_column_letter(c)}{r}")
                for r in range(min_row, max_row + 1)
                for c in range(min_col, max_col + 1)
            ]
        return self.value(sheet, cell_part)

    def _call(self, name: str, args: list[tuple], sheet: str) -> Any:
        if name == "IF":  # lazy: only the chosen branch is evaluated
            condition = self._eval(args[0], sheet)
            branch = args[1] if condition else args[2]
            return self._eval(branch, sheet)
        values = [self._eval(a, sheet) for a in args]
        if name == "SUM":
            total = 0.0
            for v in values:
                if isinstance(v, list):
                    total += sum(x for x in v if _is_number(x))
                else:
                    total += _to_number(v)
            return total
        if name == "COUNT":
            return sum(
                sum(1 for x in v if _is_number(x)) if isinstance(v, list) else int(_is_number(v))
                for v in values
            )
        if name == "COUNTIF":
            items, criterion = values
            if isinstance(criterion, str):
                return sum(1 for x in items if isinstance(x, str) and x.lower() == criterion.lower())
            return sum(1 for x in items if _is_number(x) and x == criterion)
        if name == "AND":
            flat = [x for v in values for x in (v if isinstance(v, list) else [v])]
            return all(bool(x) for x in flat)
        if name == "ABS":
            return abs(_to_number(values[0]))
        if name == "ROUND":
            return round(_to_number(values[0]), int(_to_number(values[1])))
        raise NotImplementedError(f"Function {name} is not supported by the test evaluator")


# ---------------------------------------------------------------------------
# Lookup helpers (find cells by label, so tests do not depend on row numbers)
# ---------------------------------------------------------------------------


def find_row(ws: Any, label: str, column: int = 1) -> int:
    for row in range(1, ws.max_row + 1):
        if ws.cell(row=row, column=column).value == label:
            return row
    raise AssertionError(f"Label {label!r} not found in column {column} of {ws.title}")


def labelled_value(calc: FormulaEvaluator, sheet: str, label: str, column: str = "B") -> Any:
    row = find_row(calc.wb[sheet], label)
    return calc.value(sheet, f"{column}{row}")


def row_values(calc: FormulaEvaluator, sheet: str, label: str, first_col: int, count: int) -> list[Any]:
    row = find_row(calc.wb[sheet], label)
    return [calc.value(sheet, f"{get_column_letter(first_col + i)}{row}") for i in range(count)]


def sheet_text(ws: Any) -> str:
    return " ".join(
        str(cell.value) for row in ws.iter_rows() for cell in row if isinstance(cell.value, str)
    )


def is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest.fixture(scope="module")
def inputs() -> DCFInputs:
    return DCFInputs(
        company_name="Apple Inc.",
        ticker="AAPL",
        share_price=255,
        forecast_years=5,
        revenue_growth=0.08,
        ebitda_margin=0.33,
        tax_rate=0.16,
        wacc=0.085,
        terminal_growth=0.025,
        net_debt=50000,
        shares_outstanding=15000,
    )


@pytest.fixture(scope="module")
def python_result(inputs):
    """The independent reference: the tested Python DCF engine."""
    return run_dcf(inputs)


@pytest.fixture(scope="module")
def workbook_path(tmp_path_factory, python_result):
    out_dir = tmp_path_factory.mktemp("excel_out")
    return generate_dcf_workbook(python_result, out_dir / "aapl_dcf.xlsx", generated_at=GENERATED_AT)


@pytest.fixture(scope="module")
def wb(workbook_path):
    return load_workbook(workbook_path)


@pytest.fixture(scope="module")
def calc(wb):
    return FormulaEvaluator(wb)


@pytest.fixture
def live_wb(workbook_path):
    """A fresh, editable copy of the workbook for tests that change inputs."""
    return load_workbook(workbook_path)


# ---------------------------------------------------------------------------
# 1-4. Creation, existence, opening, sheet names and order
# ---------------------------------------------------------------------------


def test_workbook_is_created_successfully(workbook_path):
    assert isinstance(workbook_path, Path)
    assert workbook_path.suffix == ".xlsx"


def test_output_file_exists_and_is_a_real_xlsx(workbook_path):
    assert workbook_path.exists() and workbook_path.is_file()
    assert workbook_path.stat().st_size > 0
    assert workbook_path.read_bytes()[:2] == b"PK"  # .xlsx files are zip archives


def test_workbook_opens_successfully(wb):
    assert len(wb.worksheets) > 0


def test_workbook_has_exactly_the_required_sheets(wb):
    assert len(wb.sheetnames) == 9
    assert set(wb.sheetnames) == set(REQUIRED_SHEETS)


def test_sheets_are_in_the_required_order(wb):
    assert wb.sheetnames == REQUIRED_SHEETS


# ---------------------------------------------------------------------------
# 5-9. Identity and headline valuation figures (read from evaluated formulas)
# ---------------------------------------------------------------------------

IDENTITY_LOCATIONS = [
    ("00_Cover", "B", {"company": "Company", "ticker": "Ticker"}),
    ("01_Assumptions", "B", {"company": "Company name", "ticker": "Ticker"}),
    ("06_Valuation_Summary", "B", {"company": "Company", "ticker": "Ticker"}),
]


@pytest.mark.parametrize("sheet,column,labels", IDENTITY_LOCATIONS)
def test_company_name_appears_correctly(calc, python_result, sheet, column, labels):
    assert labelled_value(calc, sheet, labels["company"], column) == python_result.inputs.company_name


@pytest.mark.parametrize("sheet,column,labels", IDENTITY_LOCATIONS)
def test_ticker_appears_correctly(calc, python_result, sheet, column, labels):
    assert labelled_value(calc, sheet, labels["ticker"], column) == python_result.inputs.ticker


VALUATION_LOCATIONS = {
    "enterprise_value": [
        ("00_Cover", "Enterprise value", "B"),
        ("04_DCF", "Enterprise value", "C"),
        ("06_Valuation_Summary", "Enterprise value", "B"),
    ],
    "equity_value": [
        ("00_Cover", "Equity value", "B"),
        ("04_DCF", "Equity value", "C"),
        ("06_Valuation_Summary", "Equity value", "B"),
    ],
    "implied_share_price": [
        ("00_Cover", "Implied share price", "B"),
        ("04_DCF", "Implied share price", "C"),
        ("06_Valuation_Summary", "Implied DCF share price", "B"),
    ],
}


@pytest.mark.parametrize("sheet,label,column", VALUATION_LOCATIONS["enterprise_value"])
def test_enterprise_value_appears_correctly(calc, python_result, sheet, label, column):
    assert_close(labelled_value(calc, sheet, label, column), python_result.enterprise_value,
                 f"{sheet} enterprise value")


@pytest.mark.parametrize("sheet,label,column", VALUATION_LOCATIONS["equity_value"])
def test_equity_value_appears_correctly(calc, python_result, sheet, label, column):
    assert_close(labelled_value(calc, sheet, label, column), python_result.equity_value,
                 f"{sheet} equity value")


@pytest.mark.parametrize("sheet,label,column", VALUATION_LOCATIONS["implied_share_price"])
def test_implied_share_price_appears_correctly(calc, python_result, sheet, label, column):
    assert_close(labelled_value(calc, sheet, label, column), python_result.implied_share_price,
                 f"{sheet} implied share price")


def test_cover_shows_upside_current_price_forecast_period_and_date(calc, python_result):
    sheet = "00_Cover"
    assert_close(labelled_value(calc, sheet, "Upside / (downside)"), python_result.upside_downside)
    assert_close(labelled_value(calc, sheet, "Current share price"), python_result.current_share_price)
    assert labelled_value(calc, sheet, "Forecast period") == python_result.inputs.forecast_years
    assert labelled_value(calc, sheet, "Model date") == GENERATED_AT
    assert labelled_value(calc, sheet, "Model checks") == "ALL CHECKS PASS"


def test_cover_states_this_is_an_illustrative_automated_model(wb):
    text = sheet_text(wb["00_Cover"]).lower()
    assert "dcf valuation" in text
    assert "illustrative automated dcf model" in text
    assert "controlled sample data" in text


def test_valuation_summary_contains_all_required_metrics(calc, python_result):
    sheet = "06_Valuation_Summary"
    assert_close(labelled_value(calc, sheet, "Current share price"), python_result.current_share_price)
    assert_close(labelled_value(calc, sheet, "Upside / (downside)"), python_result.upside_downside)
    assert_close(labelled_value(calc, sheet, "WACC"), python_result.inputs.wacc)
    assert_close(labelled_value(calc, sheet, "Terminal growth"), python_result.inputs.terminal_growth)


# ---------------------------------------------------------------------------
# 11. Sensitivity structure and values
# ---------------------------------------------------------------------------


def test_sensitivity_sheet_has_title_labels_and_axes(wb, calc):
    ws = wb["05_Sensitivity"]
    assert "sensitivity" in str(ws["A1"].value).lower()
    label = str(ws["A7"].value).lower()
    assert "wacc" in label and "terminal growth" in label
    growth = [calc.value("05_Sensitivity", f"{c}7") for c in GRID_COLUMNS]
    wacc = [calc.value("05_Sensitivity", f"B{r}") for r in GRID_ROWS]
    assert_all_close(growth, GROWTH_AXIS, "terminal growth axis")
    assert_all_close(wacc, WACC_AXIS, "WACC axis")


def test_sensitivity_axes_are_centred_on_the_base_case(calc, python_result):
    assert_close(calc.value("05_Sensitivity", "E7"), python_result.inputs.terminal_growth)
    assert_close(calc.value("05_Sensitivity", "B10"), python_result.inputs.wacc)


def test_sensitivity_grid_is_five_by_five_numbers(calc):
    for row in GRID_ROWS:
        for col in GRID_COLUMNS:
            value = calc.value("05_Sensitivity", f"{col}{row}")
            assert _is_number(value), f"{col}{row} is not numeric: {value!r}"


def test_sensitivity_centre_cell_is_the_base_case(calc, python_result):
    assert_close(calc.value("05_Sensitivity", "E10"), python_result.implied_share_price)


def test_every_sensitivity_cell_matches_an_independent_python_run(calc, inputs):
    for row, wacc in zip(GRID_ROWS, WACC_AXIS):
        for col, growth in zip(GRID_COLUMNS, GROWTH_AXIS):
            expected = run_dcf(replace(inputs, wacc=wacc, terminal_growth=growth)).implied_share_price
            assert_close(calc.value("05_Sensitivity", f"{col}{row}"), expected,
                         f"WACC={wacc} g={growth}")


def test_sensitivity_direction_is_financially_sensible(calc):
    grid = [[calc.value("05_Sensitivity", f"{c}{r}") for c in GRID_COLUMNS] for r in GRID_ROWS]
    for row in grid:  # higher terminal growth -> higher price
        assert row == sorted(row)
    for col in range(5):  # higher WACC -> lower price
        column = [grid[r][col] for r in range(5)]
        assert column == sorted(column, reverse=True)


def test_sensitivity_cells_are_formulas_not_hard_coded(wb):
    ws = wb["05_Sensitivity"]
    for row in GRID_ROWS:
        for col in GRID_COLUMNS:
            assert is_formula(ws[f"{col}{row}"].value), f"{col}{row} is hard-coded"
    assert "01_Assumptions" in ws["E7"].value and "01_Assumptions" in ws["B10"].value


def test_sensitivity_shows_na_where_wacc_does_not_exceed_growth(tmp_path, inputs):
    narrow = replace(inputs, wacc=0.02, terminal_growth=0.015)
    result = run_dcf(narrow)
    path = generate_dcf_workbook(result, tmp_path / "narrow.xlsx", generated_at=GENERATED_AT)
    calc = FormulaEvaluator(load_workbook(path))
    waccs = [0.01, 0.015, 0.02, 0.025, 0.03]
    growths = [0.005, 0.01, 0.015, 0.02, 0.025]
    for row, wacc in zip(GRID_ROWS, waccs):
        for col, growth in zip(GRID_COLUMNS, growths):
            value = calc.value("05_Sensitivity", f"{col}{row}")
            if wacc <= growth + 1e-12:
                assert value == "n/a", f"WACC={wacc} g={growth} should be n/a"
            else:
                expected = run_dcf(replace(narrow, wacc=wacc, terminal_growth=growth)).implied_share_price
                assert_close(value, expected, f"WACC={wacc} g={growth}")


# ---------------------------------------------------------------------------
# 12-15. Model checks and Python-vs-Excel reconciliation
# ---------------------------------------------------------------------------


def test_model_checks_sheet_lists_every_required_check(wb):
    ws = wb["07_Model_Checks"]
    names = [ws.cell(row=r, column=2).value for r in range(8, ws.max_row + 1)]
    for name in CHECK_NAMES:
        assert name in names, f"missing check: {name}"
    header = [ws.cell(row=7, column=c).value for c in range(1, 9)]
    assert header[:7] == ["#", "Check", "Python / reference", "Excel value", "Difference",
                          "Tolerance", "Result"]


@pytest.mark.parametrize("name", CHECK_NAMES)
def test_each_model_check_passes_on_a_freshly_generated_workbook(calc, name):
    sheet = "07_Model_Checks"
    row = find_row(calc.wb[sheet], name, column=2)
    assert calc.value(sheet, f"G{row}") == "PASS"


def test_overall_model_status_is_pass(calc):
    assert calc.value("07_Model_Checks", "B4") == "ALL CHECKS PASS"


def test_model_check_tolerance_is_tight(wb):
    tolerance = wb["07_Model_Checks"]["B5"].value
    assert 0 < tolerance <= 1e-6  # a large tolerance could hide real discrepancies


def _reconcile(calc, python_value, check_name):
    sheet = "07_Model_Checks"
    row = find_row(calc.wb[sheet], check_name, column=2)
    excel_value = calc.value(sheet, f"D{row}")
    assert_close(excel_value, python_value, check_name)
    assert_close(calc.value(sheet, f"E{row}"), 0.0, f"{check_name} difference")
    assert calc.value(sheet, f"G{row}") == "PASS"


def test_python_and_excel_enterprise_value_reconcile(calc, python_result):
    _reconcile(calc, python_result.enterprise_value, "Python vs Excel: Enterprise value")


def test_python_and_excel_equity_value_reconcile(calc, python_result):
    _reconcile(calc, python_result.equity_value, "Python vs Excel: Equity value")


def test_python_and_excel_implied_share_price_reconcile(calc, python_result):
    _reconcile(calc, python_result.implied_share_price, "Python vs Excel: Implied share price")


def test_model_checks_python_column_holds_engine_values(calc, python_result):
    sheet = "07_Model_Checks"
    for name, expected in [
        ("Python vs Excel: Enterprise value", python_result.enterprise_value),
        ("Python vs Excel: Equity value", python_result.equity_value),
        ("Python vs Excel: Implied share price", python_result.implied_share_price),
    ]:
        row = find_row(calc.wb[sheet], name, column=2)
        assert_close(calc.value(sheet, f"C{row}"), expected, name)


def test_a_wacc_change_makes_the_model_checks_fail_visibly(live_wb):
    # Excel formulas move with the assumption; the stored Python reference does not.
    ws = live_wb["01_Assumptions"]
    ws[f"B{find_row(ws, 'WACC')}"].value = 0.095
    calc = FormulaEvaluator(live_wb)
    sheet = "07_Model_Checks"
    row = find_row(live_wb[sheet], "Python vs Excel: Enterprise value", column=2)
    assert calc.value(sheet, f"G{row}") == "FAIL"
    assert abs(calc.value(sheet, f"E{row}")) > 1.0
    assert calc.value(sheet, "B4").startswith("CHECK FAILURE")
    # The internal bridge checks still pass: the model is consistent, just not the Python case.
    bridge = find_row(live_wb[sheet], "Enterprise value bridge", column=2)
    assert calc.value(sheet, f"G{bridge}") == "PASS"


def test_wacc_not_above_terminal_growth_fails_its_check(live_wb):
    ws = live_wb["01_Assumptions"]
    ws[f"B{find_row(ws, 'Terminal growth')}"].value = 0.09  # above the 8.5% WACC
    calc = FormulaEvaluator(live_wb)
    sheet = "07_Model_Checks"
    row = find_row(live_wb[sheet], "WACC greater than terminal growth", column=2)
    assert calc.value(sheet, f"E{row}") < 0
    assert calc.value(sheet, f"G{row}") == "FAIL"


# ---------------------------------------------------------------------------
# 16. Forecast
# ---------------------------------------------------------------------------


def test_forecast_contains_the_correct_number_of_years(wb, calc, python_result):
    ws = wb["03_Forecast"]
    expected_years = len(python_result.projections)
    assert expected_years == 5
    headers = [ws.cell(row=4, column=c).value for c in range(4, 4 + expected_years + 1)]
    assert headers[:expected_years] == [f"Year {i}" for i in range(1, expected_years + 1)]
    assert headers[expected_years] is None  # no extra forecast column
    periods = row_values(calc, "03_Forecast", "Forecast period (t)", 4, expected_years)
    assert periods == list(range(1, expected_years + 1))


FORECAST_LINES = [
    ("Revenue", "revenue"),
    ("EBITDA", "ebitda"),
    ("D&A", "depreciation_amortization"),
    ("EBIT", "ebit"),
    ("Tax", "taxes"),
    ("NOPAT", "nopat"),
    ("CapEx", "capex"),
    ("Change in NWC", "change_in_nwc"),
    ("UFCF", "ufcf"),
]


@pytest.mark.parametrize("label,attribute", FORECAST_LINES)
def test_forecast_line_items_match_the_python_engine_every_year(calc, python_result, label, attribute):
    expected = [getattr(p, attribute) for p in python_result.projections]
    actual = row_values(calc, "03_Forecast", label, 4, len(expected))
    assert_all_close(actual, expected, label)


def test_forecast_growth_margin_and_tax_rate_rows_link_to_assumptions(calc, python_result):
    n = len(python_result.projections)
    inp = python_result.inputs
    assert_all_close(row_values(calc, "03_Forecast", "Revenue growth", 4, n), [inp.revenue_growth] * n)
    assert_all_close(row_values(calc, "03_Forecast", "EBITDA margin", 4, n), [inp.ebitda_margin] * n)
    assert_all_close(row_values(calc, "03_Forecast", "Tax rate", 4, n), [inp.tax_rate] * n)


def test_forecast_uses_visible_formulas_matching_the_engine_logic(wb):
    ws = wb["03_Forecast"]
    r = {label: find_row(ws, label) for label in (
        "Revenue", "Revenue growth", "EBITDA", "EBITDA margin", "D&A", "EBIT", "Tax rate",
        "Tax", "NOPAT", "CapEx", "Change in NWC", "UFCF")}
    # Year 2 (column E) refers to year 1 (column D) and to same-column drivers.
    assert ws[f"E{r['Revenue']}"].value == f"=D{r['Revenue']}*(1+E{r['Revenue growth']})"
    assert ws[f"E{r['EBITDA']}"].value == f"=E{r['Revenue']}*E{r['EBITDA margin']}"
    assert ws[f"E{r['EBIT']}"].value == f"=E{r['EBITDA']}-E{r['D&A']}"
    assert ws[f"E{r['Tax']}"].value == f"=E{r['EBIT']}*E{r['Tax rate']}"
    assert ws[f"E{r['NOPAT']}"].value == f"=E{r['EBIT']}-E{r['Tax']}"
    assert ws[f"E{r['UFCF']}"].value == (
        f"=E{r['NOPAT']}+E{r['D&A']}-E{r['CapEx']}-E{r['Change in NWC']}"
    )
    assert ws[f"E{r['D&A']}"].value == f"=D{r['D&A']}"  # held flat
    assert "01_Assumptions" in ws[f"E{r['Revenue growth']}"].value
    assert "02_Historical" in ws[f"C{r['Revenue']}"].value  # base year links to history


@pytest.mark.parametrize("years", [1, 3, 10])
def test_workbook_adapts_to_other_forecast_lengths(tmp_path, inputs, years):
    result = run_dcf(replace(inputs, forecast_years=years))
    path = generate_dcf_workbook(result, tmp_path / f"years_{years}.xlsx", generated_at=GENERATED_AT)
    book = load_workbook(path)
    calc = FormulaEvaluator(book)
    ws = book["03_Forecast"]
    assert [ws.cell(row=4, column=c).value for c in range(4, 4 + years)] == [
        f"Year {i}" for i in range(1, years + 1)]
    assert ws.cell(row=4, column=4 + years).value is None
    assert_close(labelled_value(calc, "04_DCF", "Enterprise value", "C"), result.enterprise_value)
    assert_close(calc.value("05_Sensitivity", "E10"), result.implied_share_price)
    assert calc.value("07_Model_Checks", "B4") == "ALL CHECKS PASS"


# ---------------------------------------------------------------------------
# 17. Historical data
# ---------------------------------------------------------------------------


def test_historical_sheet_contains_the_controlled_data_exactly(calc, python_result):
    sheet = "02_Historical"
    expected = {
        "Revenue": [100000, 110000, 120000],
        "EBITDA": [25000, 28000, 31000],
        "D&A": [5000, 5500, 6000],
        "CapEx": [6000, 6500, 7000],
        "Change in NWC": [2000, 2200, 2400],
    }
    for label, values in expected.items():
        assert row_values(calc, sheet, label, 2, 3) == values, label
    hist = python_result.historical
    assert row_values(calc, sheet, "Revenue", 2, 3) == list(hist.revenue)
    assert row_values(calc, sheet, "EBITDA", 2, 3) == list(hist.ebitda)
    assert row_values(calc, sheet, "D&A", 2, 3) == list(hist.depreciation_amortization)
    assert row_values(calc, sheet, "CapEx", 2, 3) == list(hist.capex)
    assert row_values(calc, sheet, "Change in NWC", 2, 3) == list(hist.change_in_nwc)


def test_historical_periods_are_labelled_and_flagged_as_sample_data(wb):
    ws = wb["02_Historical"]
    headers = [ws.cell(row=4, column=c).value for c in range(1, 5)]
    assert headers == ["Line item", "Period 1 (oldest)", "Period 2", "Period 3 (latest)"]
    assert "controlled sample data" in str(ws["A2"].value).lower()


def test_historical_inputs_are_hard_coded_and_derived_rows_are_formulas(wb, calc):
    ws = wb["02_Historical"]
    revenue_row = find_row(ws, "Revenue")
    assert not is_formula(ws[f"B{revenue_row}"].value)
    growth_row = find_row(ws, "Revenue growth")
    assert is_formula(ws[f"C{growth_row}"].value)
    assert_close(calc.value("02_Historical", f"C{growth_row}"), 110000 / 100000 - 1)
    assert_close(calc.value("02_Historical", f"D{growth_row}"), 120000 / 110000 - 1)


def test_forecast_base_values_are_the_latest_historical_values(calc):
    sheet = "02_Historical"
    assert labelled_value(calc, sheet, "Latest revenue") == 120000
    assert labelled_value(calc, sheet, "Latest D&A") == 6000
    assert labelled_value(calc, sheet, "Latest CapEx") == 7000
    assert labelled_value(calc, sheet, "Latest change in NWC") == 2400


# ---------------------------------------------------------------------------
# 18. Assumptions
# ---------------------------------------------------------------------------

ASSUMPTION_FIELDS = [
    ("Company name", "company_name"),
    ("Ticker", "ticker"),
    ("Current share price", "share_price"),
    ("Forecast years", "forecast_years"),
    ("Revenue growth", "revenue_growth"),
    ("EBITDA margin", "ebitda_margin"),
    ("Tax rate", "tax_rate"),
    ("WACC", "wacc"),
    ("Terminal growth", "terminal_growth"),
    ("Net debt", "net_debt"),
    ("Shares outstanding", "shares_outstanding"),
]


@pytest.mark.parametrize("label,field", ASSUMPTION_FIELDS)
def test_assumptions_sheet_contains_every_input(wb, inputs, label, field):
    ws = wb["01_Assumptions"]
    cell = ws[f"B{find_row(ws, label)}"]
    expected = getattr(inputs, field)
    if isinstance(expected, str):
        assert cell.value == expected
    else:
        assert_close(cell.value, expected, label)
    assert not is_formula(cell.value)  # assumptions are the model's inputs


def test_assumptions_use_readable_number_formats(wb):
    ws = wb["01_Assumptions"]
    for label in ("Revenue growth", "EBITDA margin", "Tax rate", "WACC", "Terminal growth"):
        assert "%" in ws[f"B{find_row(ws, label)}"].number_format, label
    assert "0.00" in ws[f"B{find_row(ws, 'Current share price')}"].number_format
    assert "#,##0" in ws[f"B{find_row(ws, 'Net debt')}"].number_format


# ---------------------------------------------------------------------------
# 19. DCF bridge
# ---------------------------------------------------------------------------


def test_dcf_sheet_contains_the_bridge_in_order(wb):
    ws = wb["04_DCF"]
    order = ["PV of forecast UFCF", "PV of terminal value", "Enterprise value", "Net debt",
             "Equity value", "Shares outstanding", "Implied share price", "Current share price",
             "Upside / (downside)"]
    rows = [find_row(ws, label) for label in order]
    assert rows == sorted(rows)


def test_dcf_discounting_matches_the_python_engine_every_year(calc, python_result):
    n = len(python_result.projections)
    for label, attribute in [("UFCF", "ufcf"), ("Discount factor", "discount_factor"),
                             ("PV of UFCF", "pv_ufcf")]:
        expected = [getattr(p, attribute) for p in python_result.projections]
        assert_all_close(row_values(calc, "04_DCF", label, 3, n), expected, label)


def test_dcf_bridge_components_match_the_python_engine(calc, python_result):
    sheet = "04_DCF"
    assert_close(labelled_value(calc, sheet, "PV of forecast UFCF", "C"), python_result.sum_pv_ufcf)
    assert_close(labelled_value(calc, sheet, "Terminal value", "C"), python_result.terminal_value)
    assert_close(labelled_value(calc, sheet, "PV of terminal value", "C"), python_result.pv_terminal_value)
    assert_close(labelled_value(calc, sheet, "Net debt", "C"), python_result.net_debt)
    assert_close(labelled_value(calc, sheet, "Shares outstanding", "C"), python_result.shares_outstanding)
    assert_close(labelled_value(calc, sheet, "WACC", "C"), python_result.inputs.wacc)
    assert_close(labelled_value(calc, sheet, "Terminal growth", "C"), python_result.inputs.terminal_growth)
    assert_close(labelled_value(calc, sheet, "Upside / (downside)", "C"), python_result.upside_downside)


def test_dcf_bridge_arithmetic_holds_inside_the_workbook(calc):
    sheet = "04_DCF"
    pv_sum = labelled_value(calc, sheet, "PV of forecast UFCF", "C")
    pv_tv = labelled_value(calc, sheet, "PV of terminal value", "C")
    ev = labelled_value(calc, sheet, "Enterprise value", "C")
    net_debt = labelled_value(calc, sheet, "Net debt", "C")
    equity = labelled_value(calc, sheet, "Equity value", "C")
    shares = labelled_value(calc, sheet, "Shares outstanding", "C")
    price = labelled_value(calc, sheet, "Implied share price", "C")
    assert_close(pv_sum + pv_tv, ev, "PV of UFCF + PV of TV = EV")
    assert_close(ev - net_debt, equity, "EV - net debt = equity value")
    assert_close(equity / shares, price, "equity value / shares = implied price")


def test_dcf_bridge_cells_are_formulas_that_reference_each_other(wb):
    ws = wb["04_DCF"]
    ev_row = find_row(ws, "Enterprise value")
    sum_row = find_row(ws, "PV of forecast UFCF")
    pv_tv_row = find_row(ws, "PV of terminal value")
    assert ws[f"C{ev_row}"].value == f"=C{sum_row}+C{pv_tv_row}"
    tv_row = find_row(ws, "Terminal value")
    wacc_row = find_row(ws, "WACC")
    g_row = find_row(ws, "Terminal growth")
    final_row = find_row(ws, "Final year UFCF")
    assert ws[f"C{tv_row}"].value == (
        f"=C{final_row}*(1+C{g_row})/(C{wacc_row}-C{g_row})"
    )
    assert ws[f"C{wacc_row}"].value.startswith("='01_Assumptions'!")


def test_valuation_outputs_are_formulas_not_pasted_numbers(wb):
    for sheet, label, column in (
        VALUATION_LOCATIONS["enterprise_value"]
        + VALUATION_LOCATIONS["equity_value"]
        + VALUATION_LOCATIONS["implied_share_price"]
    ):
        ws = wb[sheet]
        value = ws[f"{column}{find_row(ws, label)}"].value
        assert is_formula(value), f"{sheet}: {label} is hard-coded ({value!r})"


# ---------------------------------------------------------------------------
# The workbook is a live model: changing any assumption changes the valuation
# exactly as the Python engine says it should.
# ---------------------------------------------------------------------------

MUTATIONS = [
    ("Revenue growth", "revenue_growth", 0.10),
    ("EBITDA margin", "ebitda_margin", 0.35),
    ("Tax rate", "tax_rate", 0.20),
    ("WACC", "wacc", 0.095),
    ("Terminal growth", "terminal_growth", 0.03),
    ("Net debt", "net_debt", 80000),
    ("Shares outstanding", "shares_outstanding", 16000),
]


@pytest.mark.parametrize("label,field,new_value", MUTATIONS)
def test_workbook_recalculates_like_the_engine_when_an_assumption_changes(
    live_wb, inputs, label, field, new_value
):
    ws = live_wb["01_Assumptions"]
    ws[f"B{find_row(ws, label)}"].value = new_value
    calc = FormulaEvaluator(live_wb)
    expected = run_dcf(replace(inputs, **{field: new_value}))
    assert_close(labelled_value(calc, "04_DCF", "Enterprise value", "C"), expected.enterprise_value, "EV")
    assert_close(labelled_value(calc, "04_DCF", "Equity value", "C"), expected.equity_value, "equity")
    assert_close(labelled_value(calc, "04_DCF", "Implied share price", "C"),
                 expected.implied_share_price, "price")
    assert_close(labelled_value(calc, "04_DCF", "Upside / (downside)", "C"),
                 expected.upside_downside, "upside")
    # The stored Python reference values no longer match, and the checks say so.
    assert calc.value("07_Model_Checks", "B4").startswith("CHECK FAILURE")


# ---------------------------------------------------------------------------
# 20. Data sources and limitations disclosure
# ---------------------------------------------------------------------------


def test_data_sources_sheet_discloses_controlled_sample_data(wb):
    text = sheet_text(wb["08_Data_Sources"]).lower()
    assert "controlled sample data" in text
    assert "not live market data" in text
    assert "demonstration and testing" in text
    assert "no market-data apis" in text


def test_data_sources_never_claims_the_data_is_live(wb):
    text = sheet_text(wb["08_Data_Sources"]).lower()
    mentions = [s for s in re.split(r"(?<=[.!?])\s+", text) if "live market data" in s]
    assert mentions, "the sheet should explicitly address live market data"
    for sentence in mentions:  # every mention must be a denial
        assert re.search(r"\b(not|no|none|never)\b", sentence), sentence


def test_data_sources_documents_methodology_and_limitations(wb):
    ws = wb["08_Data_Sources"]
    topics = [ws.cell(row=r, column=1).value for r in range(5, ws.max_row + 1)]
    assert "Historical data source / status" in topics
    assert "Model assumptions source / status" in topics
    assert "DCF methodology" in topics
    assert any(str(t).startswith("Limitation") for t in topics)
    text = sheet_text(ws).lower()
    assert "gordon growth" in text
    assert "held flat" in text
    assert "end-of-year" in text
    assert "not investment advice" in text


def test_data_sources_records_the_generation_timestamp(wb):
    ws = wb["08_Data_Sources"]
    assert ws[f"B{find_row(ws, 'Generation timestamp')}"].value == GENERATED_AT


# ---------------------------------------------------------------------------
# Output location handling
# ---------------------------------------------------------------------------


def test_directory_output_creates_a_named_file_inside_it(tmp_path, python_result):
    path = generate_dcf_workbook(python_result, tmp_path, generated_at=GENERATED_AT)
    assert path.parent == tmp_path
    assert path.name == "AAPL_DCF_20260924_103000.xlsx"
    assert path.exists()


def test_missing_output_folders_are_created(tmp_path, python_result):
    target = tmp_path / "nested" / "outputs" / "model.xlsx"
    path = generate_dcf_workbook(python_result, target, generated_at=GENERATED_AT)
    assert path == target and path.exists()


def test_non_xlsx_output_path_is_rejected(tmp_path, python_result):
    with pytest.raises(ValueError, match=r"\.xlsx"):
        generate_dcf_workbook(python_result, tmp_path / "model.csv", generated_at=GENERATED_AT)


def test_default_output_folder_is_the_project_outputs_folder():
    project_root = Path(__file__).resolve().parent.parent
    assert DEFAULT_OUTPUT_DIR == project_root / "outputs"